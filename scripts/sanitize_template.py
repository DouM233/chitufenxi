"""One-time mother-template cleanup: data areas must be white-bodied like the gold report.

Keeps title row (r1), header row(s), and light-tinted note rows; strips dark fills
from data rows and flips white text to black so it stays readable on white.
"""
import copy
import warnings
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

warnings.filterwarnings("ignore")

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "excel" / "单品客服需求分析标准母版.xlsx"

TARGET_SHEETS = {
    "新增词条": [1],
    "P0风险明细": [1, 2],
    "安全质量风险": [1, 2],
    "逐条需求明细": [1, 2],
    "同人重复记录": [1, 2],
    "待人工复核": [1, 2],
    "词条与口径": [1, 2],
    "上期基准": [1, 2],
    "原始数据": [1, 2, 3],
}

LIGHT_TINTS = {"D9EAF7", "F3F6F9", "DDEBF7", "FFF2CC"}


def is_dark(rgb):
    if not rgb or rgb in ("none", "00000000"):
        return False
    hex6 = rgb[-6:]
    if hex6.upper() in LIGHT_TINTS:
        return False
    r, g, b = (int(hex6[i : i + 2], 16) for i in (0, 2, 4))
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255 < 0.6


def clean():
    wb = load_workbook(TEMPLATE)
    for name, keep_rows in TARGET_SHEETS.items():
        ws = wb[name]
        cleaned = 0
        for row in range(1, ws.max_row + 1):
            if row in keep_rows:
                continue
            row_has_fill = False
            for cell in ws[row]:
                rgb = getattr(cell.fill.start_color, "rgb", None) if cell.fill and cell.fill.fill_type else None
                if rgb and is_dark(str(rgb)):
                    row_has_fill = True
                    break
            if not row_has_fill:
                continue
            cleaned += 1
            for cell in ws[row]:
                if cell.fill and cell.fill.fill_type:
                    fill_rgb = str(getattr(cell.fill.start_color, "rgb", "") or "")
                    if fill_rgb[-6:] not in LIGHT_TINTS:
                        cell.fill = PatternFill(fill_type=None)
                    elif fill_rgb[-6:] in LIGHT_TINTS:
                        continue
                font = cell.font
                color_rgb = str(getattr(font.color, "rgb", "") or "") if font.color else ""
                if color_rgb[-6:].upper() in ("FFFFFF", "FFF2CC", "D9EAF7", "DDEBF7"):
                    new_font = copy.copy(font)
                    from openpyxl.styles import Font

                    cell.font = Font(
                        name=new_font.name,
                        size=new_font.size,
                        bold=new_font.bold,
                        italic=new_font.italic,
                        underline=new_font.underline,
                        strike=new_font.strike,
                        vertAlign=new_font.vertAlign,
                    )
        print(f"[{name}] 清洗数据行 {cleaned} 行")
    wb.save(TEMPLATE)
    print("模板已保存:", TEMPLATE)


if __name__ == "__main__":
    clean()
