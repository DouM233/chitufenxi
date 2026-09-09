import csv
import json
import re
import shutil
import sys
import warnings
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, MergedCell

from exact_excel_writer import write_values_into_exact_template
from llm_analysis import analyze_entries_with_llm


warnings.filterwarnings("ignore", message="Unknown extension is not supported and will be removed")
warnings.filterwarnings("ignore", message="Conditional Formatting extension is not supported and will be removed")


TEMPLATE_XLSX = Path(__file__).resolve().parents[1] / "templates" / "excel" / "单品客服需求分析标准母版.xlsx"

EXPECTED_SHEETS = [
    "分析总览",
    "一级主题",
    "售前需求统计",
    "售后需求统计",
    "售前基准对比",
    "售后基准对比",
    "新增词条",
    "P0风险明细",
    "安全质量风险",
    "逐条需求明细",
    "同人重复记录",
    "待人工复核",
    "词条与口径",
    "上期基准",
    "原始数据",
]

HEADER_RE = re.compile(r"^(.+?)\s+(\d{4}-\d{2}-\d{2})\s+(\d{2}:\d{2}:\d{2})\s*$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
TIME_RE = re.compile(r"^\d{2}:\d{2}:\d{2}$")
# 网页版聊天导出常夹带非法控制字符（如 \x03），统一清洗（保留 \t \n \r）
CTRL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
URL_RE = re.compile(r"https?://\S+")
ORDER_RE = re.compile(r"(订单号|商品ID|商品号|商品编号)[:：]?\s*\d+", re.I)
SERVICE_SENDER_PATTERNS = (
    re.compile(r"^丹橘个护健康沭阳\d+$", re.I),
    re.compile(r"^jimi(?:_vender)?_\d+$", re.I),
    re.compile(r"^.+沭阳\d+$", re.I),
)
INVALID_EXACT = {
    "你好",
    "您好",
    "在吗",
    "好",
    "好的",
    "嗯",
    "嗯嗯",
    "谢谢",
    "谢谢你",
    "OK",
    "ok",
    "噢",
    "哦",
    "同意",
    "。 ",
    "。",
    "？",
    "?",
    "用户发起转人工",
    "人工",
    "人工客服",
    "转人工",
    "不用换了",
    "没有更多内容了",
    "请在客户端查看原始聊天记录，网页版无法展示",
}

LABEL_RULES = [
    ("退款退货", "售后", "售后服务", ["退款", "退货", "退了", "退掉", "退换", "申请退", "拦截", "拒收", "运费险"]),
    ("拆卸/更换配件", "售后", "使用学习", ["拆卸", "拆下", "拆装", "怎么拆", "更换", "换头", "替换", "安装", "卡扣"]),
    ("使用方法/教程", "售后", "使用学习", ["怎么用", "如何使用", "使用方法", "教程", "视频", "操作", "怎么开", "怎么弄", "怎么卷", "怎么夹", "开机", "按键"]),
    ("不加热/温度低", "售后", "加热温控", ["不热", "没温度", "温度低", "不加热", "加热慢", "温度上不去", "200度", "几乎没什么温度"]),
    ("不出雾/喷雾异常", "售后", "喷雾系统", ["不出雾", "没有雾", "没雾", "喷雾", "出雾", "水箱", "加水", "漏水"]),
    ("造型效果不好", "售后", "造型效果", ["没效果", "效果不好", "卷不出来", "卷不上", "不定型", "不持久", "夹不直", "拉不直", "不好用"]),
    ("故障/无法工作", "售后", "质量故障", ["坏了", "故障", "无法使用", "没反应", "失灵", "断电", "插电", "电源", "指示灯", "报警"]),
    ("好评/返现/补偿", "售后", "售后服务", ["好评", "返现", "返钱", "返款", "奖励金", "打款", "补偿", "红包"]),
    ("缺件/赠品/发票", "售后", "售后服务", ["缺", "少了", "赠品", "漏发", "发票", "补发"]),
    ("卷发/直发功能", "售前", "功能咨询", ["能做卷发", "可以卷发", "直卷", "直发", "卷发", "刘海", "内扣", "蓬松", "造型"]),
    ("款式/型号区别", "售前", "购买决策", ["区别", "哪款", "型号", "几款", "套餐", "五合一", "三合一", "366", "856", "选哪个"]),
    ("温度/档位/预热", "售前", "参数规格", ["多少度", "温度", "档位", "预热", "加热", "恒温"]),
    ("是否伤发/发质适用", "售前", "安全顾虑", ["伤头发", "伤发", "适合", "发质", "细软", "粗硬", "会不会伤", "烫伤"]),
    ("发货时效/物流", "售前", "交易政策", ["发货", "送到", "多久到", "物流", "次日达", "地址", "什么时候到"]),
    ("价格优惠/活动", "售前", "交易政策", ["优惠", "券", "活动", "便宜", "价格", "保价", "多少钱"]),
    ("退换/试用政策", "售前", "交易政策", ["运费险", "试用", "七天", "7天", "无理由", "质保", "保修", "售后"]),
    ("参数/尺寸/重量", "售前", "参数规格", ["重量", "尺寸", "多大", "长度", "功率", "净重", "规格"]),
    ("配件/赠品/耗材", "售前", "配件套餐", ["配件", "赠品", "梳子", "夹子", "收纳", "水", "精油"]),
]

RISK_RULES = [
    ("烫伤/安全伤害", ["烫伤", "烫到", "烫手", "烫脸", "烫头皮", "烧焦", "糊味", "冒烟"]),
    ("使用后伤发", ["头发焦", "烧头发", "伤头发", "头发坏", "断发", "掉发"]),
    ("疑似漏电/电流", ["漏电", "触电", "电流", "麻手"]),
    ("异常发热/冒烟", ["异常发热", "很烫", "冒烟", "烧了"]),
]

ACTION_MAP = {
    "使用方法/教程": "下单后和售后首触达自动发送 30 秒使用视频，包装内放步骤卡。",
    "拆卸/更换配件": "把换头/拆卸做成单独短视频，客服自动回复优先推送。",
    "不加热/温度低": "客服先排查预热时间、档位、使用方式；集中样本需品质复测温控。",
    "不出雾/喷雾异常": "补充加水、水箱、喷雾开关和清洁说明，异常样本进入售后复核。",
    "造型效果不好": "按发质、发量、温度、停留时间给标准手法，详情页明确效果边界。",
    "退款退货": "继续拆退款原因，区分不会用、效果差、故障和冲动下单，优先干预可挽回退货。",
    "卷发/直发功能": "详情页首屏增加直卷两用/适合造型示例和视频对照。",
    "款式/型号区别": "做型号对比表，明确适用人群、功能、配件和价格差异。",
    "温度/档位/预热": "详情页和说明书统一温度档位、预热时间和适用发质。",
    "是否伤发/发质适用": "明确护发卖点、适用发质和防烫注意事项，降低售前安全顾虑。",
    "发货时效/物流": "自动回复展示仓配、预计送达和偏远地区说明。",
    "价格优惠/活动": "统一优惠入口和活动口径，减少客服重复解释。",
}

PRIORITY_MAP = {
    "退款退货": "P0",
    "不加热/温度低": "P0",
    "不出雾/喷雾异常": "P0",
    "造型效果不好": "P0",
    "使用方法/教程": "P0",
    "拆卸/更换配件": "P0",
    "故障/无法工作": "P0",
    "是否伤发/发质适用": "P1",
    "款式/型号区别": "P1",
}


def set_cell(sheet, row, col, value):
    cell = sheet.cell(row, col)
    if not isinstance(cell, MergedCell):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
        cell.value = value


def write_row(sheet, row, values):
    for col, value in enumerate(values, start=1):
        set_cell(sheet, row, col, value)


def clear_values(sheet):
    for row in sheet.iter_rows():
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def ensure_rows(sheet, target_row, style_row):
    if target_row <= sheet.max_row:
        return
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[style_row].height
    for col in range(1, min(sheet.max_column, 12) + 1):
        source = sheet.cell(style_row, col)
        target = sheet.cell(target_row, col)
        if not isinstance(target, MergedCell):
            target._style = source._style
            target.number_format = source.number_format


def ensure_sheets(workbook):
    missing = [name for name in EXPECTED_SHEETS if name not in workbook.sheetnames]
    if missing:
        raise RuntimeError(f"模板缺少 Sheet：{', '.join(missing)}")


def normalize(text):
    text = URL_RE.sub("", text or "")
    return re.sub(r"\s+", " ", text).strip()


def is_service_sender(sender):
    sender_normalized = re.sub(r"\s+", "", sender or "")
    return (
        sender_normalized.startswith("jimi_")
        or "自营" in sender_normalized
        or "客服" in sender_normalized
        or any(pattern.fullmatch(sender_normalized) for pattern in SERVICE_SENDER_PATTERNS)
    )


def is_noise(text):
    if not text or text in INVALID_EXACT:
        return True
    if text.startswith("【此消息为欢迎卡片"):
        return True
    if ORDER_RE.fullmatch(text):
        return True
    if text.startswith("您选择了一个订单进行咨询") and "？" not in text and "吗" not in text:
        return True
    if text.startswith("由 服务助手 转交给") or text.startswith("由服务助手转交给"):
        return True
    return False


def is_valid_customer_message(message):
    if message.get("role") and message.get("role") != "buyer":
        return False
    if message.get("in_filter_range") is False:
        return False
    return not is_service_sender(message.get("sender", "")) and not is_noise(
        normalize(message.get("text", ""))
    )


def product_from_file(file):
    name = Path(file).stem
    return re.sub(r"聊天记录|客服聊天|\.raw|chat_fulltext|raw", "", name, flags=re.I).strip("_- ") or name


def product_from_baseline_file(file):
    name = re.sub(r"\s*\(\d+\)$", "", Path(file).stem).strip()
    match = re.match(r"(.+?)[_-](?:20\d{6})(?:[-_至](?:20\d{6}|\d{4}))?", name)
    return match.group(1).strip("_- ") if match else ""


def is_generic_product_name(value):
    return not str(value or "").strip() or str(value).strip() in {
        "未命名产品",
        "客服聊天需求分析",
    }


def parse_log_lines(product, lines, line_offset=0, start_conversation_id=0):
    messages = []
    conversation_id = start_conversation_id
    current = None
    for source_line, raw in enumerate(lines, line_offset + 1):
        line = raw.strip()
        if not line:
            continue
        if "以下为一通会话" in line:
            conversation_id += 1
            current = None
            continue
        if "会话结束" in line:
            current = None
            continue
        match = HEADER_RE.match(line)
        if match:
            sender, date, time = match.groups()
            current = {"product": product, "conversation_id": conversation_id, "sender": sender.strip(), "date": date, "time": time, "text": "", "source_line": source_line}
            messages.append(current)
        elif current:
            current["text"] = f"{current['text']}\n{line}".strip()
    return messages


def parse_log(product, file):
    lines = Path(file).read_text(encoding="utf-8", errors="replace").splitlines()
    return parse_log_lines(product, lines)


def parse_csv(product, file):
    messages = []
    with Path(file).open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        rows = list(csv.DictReader(handle))
        source_record_count = len(rows)
        for csv_line, row in enumerate(rows, start=2):
            buyer = (row.get("buyer") or row.get("买家") or row.get("客户昵称") or "").strip()
            chat_text = CTRL_CHARS_RE.sub("", row.get("chat_text") or row.get("聊天内容") or row.get("content") or row.get("message") or "")
            current_date = row.get("chat_date") or row.get("日期") or ""
            range_match = re.search(
                r"(\d{4}-\d{2}-\d{2}).*?(\d{4}-\d{2}-\d{2})",
                row.get("filter_range") or row.get("筛选范围") or "",
            )
            range_start, range_end = range_match.groups() if range_match else ("", "")
            if chat_text:
                lines = [line.strip() for line in chat_text.splitlines() if line.strip()]
                for index, line in enumerate(lines):
                    if DATE_RE.fullmatch(line):
                        current_date = line
                        continue
                    if not TIME_RE.fullmatch(line):
                        continue

                    # Full-text exports store each event as time, wrapped message,
                    # then sender. Reading it in the opposite direction assigns
                    # customer-service replies to the buyer.
                    end = index + 1
                    while (
                        end < len(lines)
                        and not TIME_RE.fullmatch(lines[end])
                        and not DATE_RE.fullmatch(lines[end])
                        and lines[end] != "没有更多内容了"
                    ):
                        end += 1
                    block = lines[index + 1 : end]
                    if len(block) < 2:
                        continue
                    sender = block[-1].strip()
                    text = "\n".join(block[:-1]).strip()
                    if not sender or not text:
                        continue
                    # role 判定：客服千牛账号必带"店铺:名字"半角冒号前缀；buyer 列（昵称）
                    # 与块内发送者（账号）常常不同名，不能只靠精确匹配，否则全部判成
                    # 客服导致"没有有效买家消息"。无冒号的按买家兜底。
                    role = "service" if ":" in sender else "buyer"
                    messages.append(
                        {
                            "product": product,
                            "conversation_id": csv_line - 1,
                            "sender": sender,
                            "role": role,
                            "date": current_date,
                            "time": line,
                            "text": text,
                            "source_line": csv_line,
                            "source_record_count": source_record_count,
                            "in_filter_range": not range_start or range_start <= current_date <= range_end,
                        }
                    )
            else:
                text = row.get("text") or row.get("内容") or ""
                if text:
                    messages.append({"product": product, "conversation_id": csv_line - 1, "sender": buyer or f"row_{csv_line}", "role": "buyer", "date": current_date, "time": "", "text": text, "source_line": csv_line, "source_record_count": source_record_count, "in_filter_range": not range_start or range_start <= current_date <= range_end})
    return messages


EXCEL_SENDER_KEYS = ("发送者", "发送人", "昵称", "账号", "买家", "客户", "用户", "sender", "buyer", "user", "name")
EXCEL_TIME_KEYS = ("时间", "日期", "date", "time")
EXCEL_TEXT_KEYS = ("内容", "消息", "聊天", "文本", "text", "message", "content", "chat")


def _excel_time_parts(value):
    """把 Excel 时间单元格归一化为 (date, time) 文本。"""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d"), value.strftime("%H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d"), ""
    if isinstance(value, (int, float)) and 20000 < value < 80000:
        base = datetime(1899, 12, 30) + timedelta(days=float(value))
        return base.strftime("%Y-%m-%d"), base.strftime("%H:%M:%S")
    text = str(value or "").strip()
    if not text:
        return "", ""
    date_match = re.search(r"(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})", text)
    date_value = f"{date_match.group(1)}-{int(date_match.group(2)):02d}-{int(date_match.group(3)):02d}" if date_match else ""
    time_match = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
    time_value = f"{int(time_match.group(1)):02d}:{time_match.group(2)}:{time_match.group(3) or '00'}" if time_match else ""
    return date_value, time_value


def _excel_header_map(header_row):
    """识别表头行并映射列；返回 (sender_idx, time_idx, text_idx) 或 None。"""
    cells = [(idx, str(cell or "").strip()) for idx, cell in enumerate(header_row)]
    if sum(1 for _, name in cells if name) < 2:
        return None
    sender_idx = time_idx = text_idx = None
    for idx, name in cells:
        low = name.lower()
        if sender_idx is None and any(key in name or key in low for key in EXCEL_SENDER_KEYS):
            sender_idx = idx
        elif time_idx is None and any(key in name or key in low for key in EXCEL_TIME_KEYS):
            time_idx = idx
        elif text_idx is None and any(key in name or key in low for key in EXCEL_TEXT_KEYS):
            text_idx = idx
    if sender_idx is None or text_idx is None:
        return None
    return sender_idx, time_idx, text_idx


def _excel_rows(file):
    """读取 xlsx/xls 第一个工作表为二维列表。"""
    ext = Path(file).suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(file, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    import xlrd

    book = xlrd.open_workbook(str(file))
    sheet = book.sheet_by_index(0)
    return [sheet.row_values(r) for r in range(sheet.nrows)]


def excel_to_csv_file(file):
    """把"一行一会话"结构的 xlsx/xls 转成 parse_csv 可直接消费的标准 csv。

    会话大文本列（单元格内含多组"账号+时间行"，网页版聊天导出同构）逐行展开为
    buyer/chat_date/chat_text 三列 csv，之后完全复用已验证的 csv 解析路径；
    转换出的 csv 保留在任务目录，可用于人工核对或重复上传。
    返回 csv 路径；非会话结构（逐行消息/无表头）返回 None，调用方走原 parse_excel。
    """
    rows = _excel_rows(file)
    if not rows:
        return None
    header = None
    data_start = 0
    for idx, row in enumerate(rows[:5]):
        mapped = _excel_header_map(row)
        if mapped:
            header = mapped
            data_start = idx + 1
            break
    if not header:
        return None
    sender_idx, time_idx, text_idx = header
    sample_texts = []
    for row in rows[data_start:]:
        value = str(row[text_idx]).strip() if text_idx < len(row) and row[text_idx] is not None else ""
        if value:
            sample_texts.append(value)
        if len(sample_texts) >= 5:
            break
    session_like = sum(1 for value in sample_texts if len(re.findall(r"\d{2}:\d{2}:\d{2}", value)) >= 2)
    if not sample_texts or session_like < max(1, len(sample_texts) // 2):
        return None
    csv_path = Path(file).parent / (Path(file).stem + ".converted.csv")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["buyer", "chat_date", "chat_text"])
        for row in rows[data_start:]:
            sender = str(row[sender_idx]).strip() if sender_idx < len(row) and row[sender_idx] is not None else ""
            text = str(row[text_idx]).strip() if text_idx < len(row) and row[text_idx] is not None else ""
            if not text:
                continue
            date_value = ""
            if time_idx is not None and time_idx < len(row):
                date_value, _time_value = _excel_time_parts(row[time_idx])
            writer.writerow([sender, date_value, CTRL_CHARS_RE.sub("", text)])
    return csv_path


def parse_excel(product, file):
    rows = _excel_rows(file)
    if not rows:
        return []
    header = None
    data_start = 0
    for idx, row in enumerate(rows[:5]):
        mapped = _excel_header_map(row)
        if mapped:
            header = mapped
            data_start = idx + 1
            break
    messages = []
    if header:
        sender_idx, time_idx, text_idx = header
        for source_line, row in enumerate(rows[data_start:], data_start + 1):
            text = str(row[text_idx]).strip() if text_idx < len(row) and row[text_idx] is not None else ""
            if not text:
                continue
            sender = str(row[sender_idx]).strip() if sender_idx < len(row) and row[sender_idx] is not None else ""
            date_value, time_value = ("", "")
            if time_idx is not None and time_idx < len(row):
                date_value, time_value = _excel_time_parts(row[time_idx])
            messages.append({
                "product": product,
                "conversation_id": source_line - data_start - 1,
                "sender": sender or f"row_{source_line}",
                "date": date_value,
                "time": time_value,
                "text": text,
                "source_line": source_line,
            })
        return messages
    # 无表头兜底：把含"头行+正文"的大单元格按 log 文本格式解析
    conversation_id = 0
    for source_line, row in enumerate(rows, 1):
        for cell in row:
            block = str(cell or "").strip()
            if not block:
                continue
            block_lines = block.splitlines()
            if not HEADER_RE.match(block_lines[0]):
                continue
            parsed = parse_log_lines(product, block_lines, source_line, conversation_id)
            if parsed:
                conversation_id = parsed[-1]["conversation_id"] + 1
                messages.extend(parsed)
    return messages


def parse_messages(product, file):
    ext = Path(file).suffix.lower()
    if ext == ".csv":
        return parse_csv(product, file)
    if ext in (".xlsx", ".xls"):
        # 会话大文本结构的 Excel 先转成标准 csv，完全复用已验证的 csv 解析路径
        converted = excel_to_csv_file(file)
        if converted:
            return parse_csv(product, converted)
        return parse_excel(product, file)
    return parse_log(product, file)


def valid_customer_messages(messages):
    result = []
    for msg in messages:
        text = normalize(msg.get("text", ""))
        if not is_valid_customer_message(msg):
            continue
        copy = dict(msg)
        copy["text"] = text
        result.append(copy)
    return result


def classify_message(text):
    labels = []
    for label, stage, theme, keywords in LABEL_RULES:
        if any(keyword in text for keyword in keywords):
            labels.append((stage, theme, label))
    seen = set()
    unique = []
    for stage, theme, label in labels:
        key = (stage, label)
        if key not in seen:
            seen.add(key)
            unique.append((stage, theme, label))
    return unique


def risk_items(text):
    return [label for label, keywords in RISK_RULES if any(keyword in text for keyword in keywords)]


def label_definition(label):
    definitions = {
        "使用方法/教程": "已购买或收到后咨询开机、安装、调档、换头、卷发/直发具体操作。",
        "拆卸/更换配件": "咨询加热管、直发梳、卷筒、梳齿等配件拆装、更换和安装方法。",
        "不加热/温度低": "实际使用后反馈温度低、不加热、预热无效或达不到标称温度。",
        "不出雾/喷雾异常": "实际使用后反馈喷雾不出、雾量异常、水箱/加水相关问题。",
        "造型效果不好": "实际使用后反馈卷不出、夹不直、不定型、效果不明显或不持久。",
        "退款退货": "实际提出退款、退货、拒收、拦截、退换货或运费险相关售后处理。",
        "卷发/直发功能": "购买前咨询是否能卷发、直发、刘海、内扣、蓬松等造型。",
        "款式/型号区别": "购买前咨询不同型号、套餐、功能、价格和适用场景差异。",
        "温度/档位/预热": "购买前咨询温度范围、档位、预热时间和加热能力。",
        "是否伤发/发质适用": "购买前咨询是否伤发、是否适合自己的发质，以及安全顾虑。",
    }
    return definitions.get(label, f"{label}相关咨询或反馈。")


def analyze_product(product, file):
    raw = parse_messages(product, file)
    valid = valid_customer_messages(raw)
    detail = []
    duplicates = []
    reviews = []
    risks = []
    retained = {}
    raw_hits = Counter()
    buyers_by_label = defaultdict(set)
    evidence_by_label = defaultdict(list)

    for msg in valid:
        labels = classify_message(msg["text"])
        if not labels:
            if any(token in msg["text"] for token in ("不", "没", "退", "坏", "怎么", "为什么", "吗")):
                reviews.append({**msg, "reason": "存在需求语气但未命中明确词条，需人工判断。"})
            continue
        for stage, theme, label in labels:
            raw_hits[(stage, label)] += 1
            key = (msg["sender"], stage, label)
            if key in retained:
                duplicates.append({**msg, "stage": stage, "label": label, "kept_text": retained[key]["text"]})
                continue
            retained[key] = msg
            buyers_by_label[(stage, theme, label)].add(msg["sender"])
            evidence_by_label[(stage, theme, label)].append(msg)
            detail.append({**msg, "stage": stage, "theme": theme, "label": label, "all_text": msg["text"], "raw_hit": raw_hits[(stage, label)]})
        for risk_label in risk_items(msg["text"]):
            risks.append({**msg, "risk_label": risk_label, "action": "进入安全/质量二次复核：确认是否为本人实际发生，记录档位、使用时长、批次和处理结果。"})

    stats = []
    for (stage, theme, label), buyers in buyers_by_label.items():
        evidences = evidence_by_label[(stage, theme, label)]
        stats.append({
            "product": product,
            "stage": stage,
            "theme": theme,
            "label": label,
            "count": len(buyers),
            "raw_hits": raw_hits[(stage, label)],
            "duplicates": sum(1 for item in duplicates if item["stage"] == stage and item["label"] == label),
            "priority": PRIORITY_MAP.get(label, "P1" if len(buyers) >= 5 else "P2"),
            "definition": label_definition(label),
            "buyers": "、".join(sorted(buyers)[:10]),
            "quote": "｜".join(item["text"] for item in evidences[:5]),
            "action": ACTION_MAP.get(label, "沉淀为 V1 词条，后续观察人数、占比和客服消耗。"),
        })

    dates = [msg["date"] for msg in valid if msg.get("date")]
    return {
        "summary": {
            "product": product,
            "raw_messages": max([msg.get("source_record_count", 0) for msg in raw] or [0]) or len(raw),
            "valid_messages": len(valid),
            "conversations": len({msg.get("conversation_id", 0) for msg in raw}),
            "buyers": len(set(msg["sender"] for msg in valid)),
            "start_date": min(dates) if dates else "",
            "end_date": max(dates) if dates else "",
            "demand_count": len(detail),
            "presale_count": sum(1 for item in detail if item["stage"] == "售前"),
            "aftersale_count": sum(1 for item in detail if item["stage"] == "售后"),
            "duplicate_count": len(duplicates),
            "risk_count": len(risks),
            "review_count": len(reviews),
        },
        "stats": stats,
        "detail": detail,
        "duplicates": duplicates,
        "risks": risks,
        "reviews": reviews,
        "valid_messages": valid,
    }


def sorted_stats(rows):
    return sorted(rows, key=lambda item: (-item["count"], item["label"]))


def stage_stats(results, stage):
    return sorted_stats([item for result in results for item in result["stats"] if item["stage"] == stage])


def aggregate_stats(results, stage):
    bucket = {}
    for item in stage_stats(results, stage):
        key = item["label"]
        if key not in bucket:
            bucket[key] = {**item, "count": 0, "raw_hits": 0, "duplicates": 0, "buyers_list": [], "quotes": []}
        bucket[key]["count"] += item["count"]
        bucket[key]["raw_hits"] += item["raw_hits"]
        bucket[key]["duplicates"] += item["duplicates"]
        bucket[key]["buyers_list"].append(item.get("buyers", ""))
        bucket[key]["quotes"].append(item.get("quote", ""))
    rows = []
    for row in bucket.values():
        row["product"] = "合计"
        row["buyers"] = "、".join(x for x in row["buyers_list"] if x)[:300]
        row["quote"] = "｜".join(x for x in row["quotes"] if x)[:500]
        rows.append(row)
    return sorted_stats(rows)


def aggregate_theme_rows(results):
    bucket = {}
    total = sum(result["summary"]["demand_count"] for result in results) or 1
    for result in results:
        for row in result["stats"]:
            theme = row.get("theme") or "其他"
            item = bucket.setdefault(theme, {"count": 0, "labels": [], "actions": []})
            item["count"] += row["count"]
            item["labels"].append(f"{row['stage']}:{row['label']}({row['count']})")
            if row.get("action"):
                item["actions"].append(row["action"])
    rows = [["一级主题", "去重需求数", "需求占比", "覆盖词条", "业务判断", "建议动作"]]
    ordered = sorted(bucket.items(), key=lambda pair: (-pair[1]["count"], pair[0]))
    for rank, (theme, item) in enumerate(ordered, 1):
        labels = "；".join(item["labels"][:8])
        judgment = f"该主题共 {item['count']} 次去重需求，位列主题第 {rank}，主要由{labels}构成。"
        rows.append(
            [
                theme,
                item["count"],
                item["count"] / total,
                labels,
                judgment,
                item["actions"][0] if item["actions"] else "持续跟踪该主题的需求人数、占比和处理结果。",
            ]
        )
    return rows


def infer_theme_from_label(label):
    text = str(label or "")
    theme_patterns = (
        ("质量安全", r"故障|异常|失效|不热|不转|漏水|漏电|破损|开裂|异响|伤害|烫伤|冒烟|二手"),
        ("使用学习", r"使用|教程|安装|组装|操作|清洗|排水|加水|拆装|收纳|维护"),
        ("参数规格", r"尺寸|桶深|重量|功率|耗电|温度|速度|材质|容量|续航|接口"),
        ("功能配置", r"功能|按摩|杀菌|红光|喷淋|冲浪|加热|恒温|配件|赠品"),
        ("型号选购", r"款式|版本|型号|区别|选择|适用|推荐|颜色"),
        ("交易履约", r"价格|优惠|国补|价保|发票|正品|自营|发货|物流|地址|预约|到货"),
        ("售后服务", r"退货|退款|换货|退换|运费|质保|保修|维修|补发|错发|少件"),
        ("体验效果", r"效果|舒适|疼|力度|噪音|气味|药包|泡脚|水温"),
    )
    for theme, pattern in theme_patterns:
        if re.search(pattern, text):
            return theme
    return "其他"


def read_baseline_data(file_path):
    empty = {
        "售前": {"rows": {}, "total": 0, "days": 0},
        "售后": {"rows": {}, "total": 0, "days": 0},
        "context": "",
        "taxonomy": [],
        "products": {},
        "start_date": "",
        "end_date": "",
    }
    if not file_path or not Path(file_path).exists():
        return empty
    workbook = load_workbook(file_path, data_only=True, read_only=True)

    def rows_for(sheet_name):
        if sheet_name not in workbook.sheetnames:
            return []
        return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]

    def number(value, default=0):
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return default

    def stage_names(value):
        text = str(value or "")
        return [stage for stage in ("售前", "售后") if stage in text]

    def sku_tokens(value):
        return list(
            dict.fromkeys(
                token.upper()
                for token in re.findall(r"(?i)(?:\d+[A-Z][A-Z0-9]*|[A-Z]+\d+[A-Z0-9]*)", str(value or ""))
            )
        )

    def put_stat(target, stage, label, count, share, daily, priority="P2", definition="", action=""):
        target[stage]["rows"][str(label)] = {
            "count": int(number(count)),
            "share": number(share),
            "daily": number(daily),
            "priority": str(priority or "P2"),
            "definition": str(definition or ""),
            "action": str(action or ""),
        }

    # Single-product standard report.
    for stage, sheet_name in (("售前", "售前需求统计"), ("售后", "售后需求统计")):
        rows = rows_for(sheet_name)
        header_index = None
        headers = []
        for index, row in enumerate(rows):
            first = row[0] if row else None
            row_headers = [str(value or "").strip() for value in row]
            if any(value in ("需求词条", "新增词条", "售前需求词条", "售后需求词条") for value in row_headers):
                header_index = index
                headers = row_headers
            if first == "正式需求总数":
                empty[stage]["total"] = int(number(row[1] if len(row) > 1 else 0))
                if len(row) > 3 and row[2] == "统计天数":
                    empty[stage]["days"] = int(number(row[3]))
        if header_index is not None:
            def stat_col(*names, default=None):
                for name in names:
                    if name in headers:
                        return headers.index(name)
                return default

            rank_col = stat_col("排名", default=0)
            label_col = stat_col("需求词条", "新增词条", "售前需求词条", "售后需求词条", default=1)
            count_col = stat_col("同人去重人数", "同人去重需求数", "去重需求数", "本期人数", default=2)
            share_col = stat_col("占比", "本期占比", "占售前", "占售后", default=3)
            daily_col = stat_col("日均人数", "日均", "本期日均")
            priority_col = stat_col("优先级", default=7)
            definition_col = stat_col("严格定义", default=8)
            action_col = stat_col("建议动作", default=11)
            for row in rows[header_index + 1 :]:
                label = row[label_col] if label_col < len(row) else None
                rank = row[rank_col] if rank_col < len(row) else None
                if not label or not isinstance(rank, (int, float)):
                    continue
                put_stat(
                    empty,
                    stage,
                    label,
                    row[count_col] if count_col < len(row) else 0,
                    row[share_col] if share_col < len(row) else 0,
                    row[daily_col] if daily_col is not None and daily_col < len(row) else 0,
                    row[priority_col] if priority_col < len(row) else "P2",
                    row[definition_col] if definition_col < len(row) else "",
                    row[action_col] if action_col < len(row) else "",
                )

    # Multi-product formal report: parse the group board.
    current_stage = None
    for row in rows_for("商品组总盘"):
        first = str(row[0] or "") if row else ""
        if "商品组售前需求" in first:
            current_stage = "售前"
            continue
        if "商品组售后需求" in first:
            current_stage = "售后"
            continue
        if current_stage and row and isinstance(row[0], (int, float)) and len(row) > 2 and row[1]:
            put_stat(
                empty,
                current_stage,
                row[1],
                row[2],
                row[3] if len(row) > 3 else 0,
                0,
                row[4] if len(row) > 4 else "P2",
                row[5] if len(row) > 5 else "",
                row[6] if len(row) > 6 else "",
            )

    # Period and dedupe context from the formal scope sheet.
    for row in rows_for("口径说明"):
        if not row:
            continue
        if str(row[0] or "") == "统计时间":
            match = re.search(r"(\d{4}-\d{2}-\d{2}).*?(\d{4}-\d{2}-\d{2})", str(row[1] if len(row) > 1 else ""))
            if match:
                empty["start_date"], empty["end_date"] = match.groups()
                start = datetime.strptime(empty["start_date"], "%Y-%m-%d").date()
                end = datetime.strptime(empty["end_date"], "%Y-%m-%d").date()
                days = (end - start).days + 1
                empty["售前"]["days"] = empty["售前"]["days"] or days
                empty["售后"]["days"] = empty["售后"]["days"] or days
    if not empty["start_date"]:
        overview_text = " ".join(
            str(value or "")
            for row in rows_for("分析总览")[:6]
            for value in row[:4]
        )
        match = re.search(r"(\d{4}-\d{2}-\d{2}).*?(\d{4}-\d{2}-\d{2})", overview_text)
        if match:
            empty["start_date"], empty["end_date"] = match.groups()
            start = datetime.strptime(empty["start_date"], "%Y-%m-%d").date()
            end = datetime.strptime(empty["end_date"], "%Y-%m-%d").date()
            days = (end - start).days + 1
            empty["售前"]["days"] = empty["售前"]["days"] or days
            empty["售后"]["days"] = empty["售后"]["days"] or days
    if not empty["start_date"]:
        name_match = re.search(r"(20\d{6})[-_至](?:(20\d{6})|(\d{4}))", Path(file_path).stem)
        if name_match:
            start_raw = name_match.group(1)
            end_raw = name_match.group(2) or start_raw[:4] + name_match.group(3)
            empty["start_date"] = datetime.strptime(start_raw, "%Y%m%d").strftime("%Y-%m-%d")
            empty["end_date"] = datetime.strptime(end_raw, "%Y%m%d").strftime("%Y-%m-%d")

    taxonomy = []
    taxonomy_seen = set()
    dictionary_rows = rows_for("V1词条字典")
    if dictionary_rows:
        for row in dictionary_rows[1:]:
            label = str(row[0] or "").strip() if row else ""
            stages = stage_names(row[2] if len(row) > 2 else "")
            if not label or not stages or label in taxonomy_seen:
                continue
            taxonomy_seen.add(label)
            taxonomy.append(
                {
                    "label": label,
                    "stage": stages[0],
                    "stages": stages,
                    "product_prefixes": sku_tokens(row[1] if len(row) > 1 else ""),
                    "theme": infer_theme_from_label(label),
                    "definition": str(row[5] or "") if len(row) > 5 else "",
                    "priority": str(row[3] or "P2") if len(row) > 3 else "P2",
                    "action": str(row[6] or "") if len(row) > 6 else "",
                }
            )
    else:
        standard_rows = rows_for("词条与口径")
        standard_header_index = next(
            (
                index
                for index, row in enumerate(standard_rows)
                if any(str(value or "").strip() in ("需求词条", "V1需求词条") for value in row)
            ),
            None,
        )
        headers = [
            str(value or "").strip()
            for value in (standard_rows[standard_header_index] if standard_header_index is not None else [])
        ]

        def header_index(*names, default=None):
            for name in names:
                if name in headers:
                    return headers.index(name)
            return default

        label_col = header_index("需求词条", "V1需求词条", default=0)
        stage_col = header_index("本期阶段", "阶段", default=1)
        theme_col = header_index("一级主题")
        definition_col = header_index("严格定义", default=2)
        priority_col = header_index("优先级", default=3)
        action_col = header_index("建议动作", default=7)
        for row in standard_rows[(standard_header_index + 1) if standard_header_index is not None else 1 :]:
            label = str(row[label_col] or "").strip() if row and label_col < len(row) else ""
            stages = stage_names(row[stage_col] if stage_col < len(row) else "")
            if not label or not stages or label in taxonomy_seen:
                continue
            taxonomy_seen.add(label)
            taxonomy.append(
                {
                    "label": label,
                    "stage": stages[0],
                    "stages": stages,
                    "product_prefixes": [],
                    "theme": str(row[theme_col] or "").strip() if theme_col is not None and theme_col < len(row) and row[theme_col] else infer_theme_from_label(label),
                    "definition": str(row[definition_col] or "") if definition_col < len(row) else "",
                    "priority": str(row[priority_col] or "P2") if priority_col < len(row) else "P2",
                    "action": str(row[action_col] or "") if action_col < len(row) else "",
                }
            )

    # Per-SKU baseline rows are used for product-level comparisons.
    ignored_sheets = set(EXPECTED_SHEETS) | {"四SKU概览", "商品组总盘", "质量安全风险", "每日趋势", "逐条需求明细", "同人重复记录", "待人工复核", "V1词条字典", "口径说明", "有效买家消息"}
    for sheet_name in workbook.sheetnames:
        if sheet_name in ignored_sheets:
            continue
        rows = rows_for(sheet_name)
        if not any(row and str(row[0] or "") in ("售前V1", "售后V1") for row in rows):
            continue
        product_name = str(rows[0][0] or sheet_name).split("｜", 1)[0]
        product_data = {
            "售前": {"rows": {}, "total": 0, "days": empty["售前"]["days"]},
            "售后": {"rows": {}, "total": 0, "days": empty["售后"]["days"]},
        }
        stage = None
        for row in rows:
            first = str(row[0] or "") if row else ""
            if first in ("售前V1", "售后V1"):
                stage = first[:2]
                continue
            if stage and row and isinstance(row[0], (int, float)) and len(row) > 2 and row[1]:
                put_stat(
                    product_data,
                    stage,
                    row[1],
                    row[2],
                    row[3] if len(row) > 3 else 0,
                    row[4] if len(row) > 4 else 0,
                    row[7] if len(row) > 7 else "P2",
                    row[8] if len(row) > 8 else "",
                    row[11] if len(row) > 11 else "",
                )
        for stage in ("售前", "售后"):
            product_data[stage]["total"] = sum(item["count"] for item in product_data[stage]["rows"].values())
        empty["products"][product_name] = product_data

    if empty["products"]:
        # Current multi-product reports sum independently deduped SKU counts, so
        # rebuild the baseline group rows from the same per-SKU unit.
        for stage in ("售前", "售后"):
            combined = {}
            for product_data in empty["products"].values():
                for label, item in product_data[stage]["rows"].items():
                    target = combined.setdefault(
                        label,
                        {
                            **empty[stage]["rows"].get(label, item),
                            "count": 0,
                        },
                    )
                    target["count"] += item["count"]
            total = sum(item["count"] for item in combined.values())
            for item in combined.values():
                item["share"] = item["count"] / total if total else 0
                item["daily"] = item["count"] / empty[stage]["days"] if empty[stage]["days"] else 0
            empty[stage]["rows"] = combined
            empty[stage]["total"] = total

    for stage in ("售前", "售后"):
        if not empty[stage]["total"]:
            empty[stage]["total"] = sum(item["count"] for item in empty[stage]["rows"].values())
        if empty[stage]["days"]:
            for item in empty[stage]["rows"].values():
                if not item["daily"]:
                    item["daily"] = item["count"] / empty[stage]["days"]

    empty["taxonomy"] = taxonomy
    empty["context"] = json.dumps(taxonomy, ensure_ascii=False)
    workbook.close()
    return empty


MANUAL_BASELINE_TOTAL_KEYS = ("正式需求总数", "基准总需求", "需求覆盖买家", "本期总需求")


def parse_manual_baseline(text, days=None, period=None):
    """Parse operator-pasted last-period rows into read_baseline_data's structure.

    Deterministic on purpose: baseline numbers must never be transcribed by the
    model. Tolerant per-row format: [排名] 词条 人数 [占比] [日均] [优先级] [定义...]；
    售前/售后标题行分组。Excel 复制的 TSV、多空格、单空格、逗号分隔都接受，
    最少只需「词条 + 人数」两列。
    """
    result = {
        "售前": {"rows": {}, "total": 0, "days": 0},
        "售后": {"rows": {}, "total": 0, "days": 0},
        "context": "",
        "taxonomy": [],
        "products": {},
        "start_date": "",
        "end_date": "",
    }
    full_text = str(text or "")
    if not full_text.strip():
        raise RuntimeError("上期基准录入内容为空。请在输入框粘贴售前/售后词条行，或改用上传基准 Excel。")

    parsed_days = 0
    if days is not None:
        try:
            parsed_days = int(days)
        except (TypeError, ValueError):
            parsed_days = 0
        if parsed_days < 0:
            parsed_days = 0
    period_match = re.search(
        r"(\d{4}-\d{2}-\d{2})\s*[至到~～—\-]{1,2}\s*(\d{4}-\d{2}-\d{2})",
        f"{period or ''}\n{full_text}",
    )
    if period_match:
        result["start_date"], result["end_date"] = period_match.groups()
        if not parsed_days:
            parsed_days = (
                datetime.strptime(result["end_date"], "%Y-%m-%d").date()
                - datetime.strptime(result["start_date"], "%Y-%m-%d").date()
            ).days + 1

    def split_tokens(line):
        # 依次尝试 Tab、多空格/逗号、单空格三种切法，取切出 token 最多的方案，
        # 避免「开头双空格 + 后续单空格」这类混合格式被提前截断。
        candidates = (
            [token.strip() for token in line.split("\t") if token.strip()],
            [token.strip() for token in re.split(r"\s{2,}|[，,｜|]", line) if token.strip()],
            line.split(),
        )
        best = []
        for tokens in candidates:
            if len(tokens) > len(best):
                best = tokens
        return best if len(best) >= 2 else []

    def is_number(token):
        return bool(re.fullmatch(r"\d+(?:,\d{3})*(?:\.\d+)?", token))

    def is_percent(token):
        return bool(re.fullmatch(r"\d+(?:\.\d+)?%", token))

    current_stage = None
    row_counts = {"售前": 0, "售后": 0}
    share_sums = {"售前": 0.0, "售后": 0.0}
    orphan_lines = []

    for raw_line in full_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        compact = re.sub(r"[\s【\[\]（）()：:，,。.、_\-]+", "", line)
        if compact in ("售前", "售后", "售前需求统计", "售后需求统计", "售前V1", "售后V1", "售前需求", "售后需求"):
            current_stage = "售前" if "售前" in compact else "售后"
            continue
        if any(key in line for key in ("统计时间", "统计周期", "分析周期")):
            continue
        days_match = re.search(r"统计天数[:：]?\s*(\d+)", line)
        if days_match:
            if not parsed_days:
                parsed_days = int(days_match.group(1))
            continue
        if any(key in line for key in MANUAL_BASELINE_TOTAL_KEYS):
            continue
        if not re.search(r"\d", line.replace("%", "")):
            # 表头、说明文字等不含数字的行直接跳过；数据行必须带人数。
            continue

        tokens = split_tokens(line)
        if len(tokens) >= 3 and re.fullmatch(r"\d+[.、]?", tokens[0]):
            tokens = tokens[1:]
        label_index = None
        for index, token in enumerate(tokens):
            if not is_number(token) and not is_percent(token):
                label_index = index
                break
        if label_index is None:
            continue
        label = tokens[label_index].strip()
        if not label or len(label) > 40:
            continue

        count = None
        share = 0.0
        priority = ""
        definition_tokens = []
        cursor = label_index + 1
        while cursor < len(tokens):
            token = tokens[cursor]
            if count is None and re.fullmatch(r"\d+(?:,\d{3})*", token):
                count = int(token.replace(",", ""))
            elif not share and is_percent(token):
                share = float(token[:-1]) / 100
            elif not share and is_number(token) and 0 < float(token) <= 1:
                share = float(token)
            elif count is not None and is_number(token):
                pass  # 日均等后续数字列，不属于定义
            elif not priority and token.upper() in ("P0", "P1", "P2"):
                priority = token.upper()
            else:
                definition_tokens = tokens[cursor:]
                break
            cursor += 1
        if count is None:
            continue
        if current_stage is None:
            orphan_lines.append(line)
            continue

        stage_rows = result[current_stage]["rows"]
        if label in stage_rows:
            raise RuntimeError(
                f"上期基准里「{label}」在{current_stage}出现多次，请只保留一行。"
            )
        stage_rows[label] = {
            "count": count,
            "share": share,
            "daily": 0,
            "priority": priority or "P2",
            "definition": re.sub(r"\s+", " ", " ".join(definition_tokens)).strip()[:300],
            "action": "",
        }
        row_counts[current_stage] += 1
        if share:
            share_sums[current_stage] += share

    if orphan_lines:
        raise RuntimeError(
            "粘贴内容里有带数字的行缺少「售前/售后」分组标题，无法确定归属："
            + "；".join(orphan_lines[:3])
        )
    if not row_counts["售前"] and not row_counts["售后"]:
        raise RuntimeError(
            "没有从录入内容识别出任何「词条 + 人数」行。请每行一条，"
            "先用「售前」「售后」标题分组，例如：使用方法/教程 79 34.1%"
        )
    for stage in ("售前", "售后"):
        if share_sums[stage] > 1.3:
            raise RuntimeError(
                f"{stage}词条占比合计约 {share_sums[stage]:.0%}，超过 130%，"
                "请检查是否误复制了对比表中本期和基准两列。"
            )
        result[stage]["total"] = sum(item["count"] for item in result[stage]["rows"].values())
        result[stage]["days"] = parsed_days
        if parsed_days:
            for item in result[stage]["rows"].values():
                item["daily"] = item["count"] / parsed_days

    taxonomy_by_label = {}
    for stage in ("售前", "售后"):
        for label, row in result[stage]["rows"].items():
            item = taxonomy_by_label.get(label)
            if item:
                if stage not in item["stages"]:
                    item["stages"].append(stage)
                continue
            taxonomy_by_label[label] = {
                "label": label,
                "stage": stage,
                "stages": [stage],
                "product_prefixes": [],
                "theme": infer_theme_from_label(label),
                "definition": row["definition"] or f"{label}相关的明确买家需求。",
                "priority": row["priority"],
                "action": row["action"] or "持续观察人数、占比和客服承接消耗。",
            }
    result["taxonomy"] = list(taxonomy_by_label.values())
    result["context"] = json.dumps(result["taxonomy"], ensure_ascii=False)
    return result


def write_rows(sheet, start_row, rows, style_row=None):
    style_row = style_row or start_row
    for offset, row in enumerate(rows):
        target = start_row + offset
        ensure_rows(sheet, target, style_row)
        write_row(sheet, target, row)


def write_overview(workbook, data, results, summary, baseline_data=None):
    sheet = workbook["分析总览"]
    clear_values(sheet)
    product = data.get("product_name") or "未命名产品"
    all_summary = [item["summary"] for item in results]
    total_raw = sum(item["raw_messages"] for item in all_summary)
    total_valid = sum(item["valid_messages"] for item in all_summary)
    total_buyers = len(
        {
            message["sender"]
            for result in results
            for message in result["valid_messages"]
        }
    )
    total_demand = sum(item["demand_count"] for item in all_summary)
    total_presale = sum(item["presale_count"] for item in all_summary)
    total_aftersale = sum(item["aftersale_count"] for item in all_summary)
    total_duplicates = sum(item["duplicate_count"] for item in all_summary)
    total_risks = sum(item["risk_count"] for item in all_summary)
    total_reviews = sum(item["review_count"] for item in all_summary)
    presale = aggregate_stats(results, "售前")
    aftersale = aggregate_stats(results, "售后")
    baseline_data = baseline_data or {"售前": {"rows": {}}, "售后": {"rows": {}}}

    set_cell(sheet, 1, 1, f"{product}｜客服聊天需求分析报告")
    set_cell(sheet, 2, 1, "本报告由模型读取本次聊天后重新建立词条并逐条语义分析；Excel 母版仅用于表格格式。")
    write_row(sheet, 4, ["数据规模", None, "消息规模", None, "需求结果", None, "重复治理", None, "安全质量", None, "交付文件", None])
    write_row(sheet, 5, ["原始记录", "有效买家", "有效买家消息", "正式需求", "售前需求", "售后需求", "重复剔除", "需求覆盖买家", "安全质量风险", "待复核", "Excel状态", "分析版本"])
    write_row(sheet, 6, [total_raw, total_buyers, total_valid, total_demand, total_presale, total_aftersale, total_duplicates, total_buyers, total_risks, total_reviews, "已生成", "AI-V1"])
    write_row(sheet, 8, ["排名", "售前词条", "本期人数", "本期占比", "基准占比", "占比变化", None, "排名", "售后词条", "本期人数", "基准占比", "占比变化"])
    for idx in range(12):
        pre = presale[idx] if idx < len(presale) else None
        aft = aftersale[idx] if idx < len(aftersale) else None
        pre_base = baseline_data["售前"]["rows"].get(pre["label"], {}) if pre else {}
        aft_base = baseline_data["售后"]["rows"].get(aft["label"], {}) if aft else {}
        pre_share = pre["count"] / total_presale if pre and total_presale else None
        aft_share = aft["count"] / total_aftersale if aft and total_aftersale else None
        write_row(sheet, 9 + idx, [
            idx + 1 if pre else None,
            pre["label"] if pre else None,
            pre["count"] if pre else None,
            pre_share,
            pre_base.get("share") if pre_base else None,
            pre_share - pre_base.get("share", 0) if pre_share is not None and pre_base else None,
            None,
            idx + 1 if aft else None,
            aft["label"] if aft else None,
            aft["count"] if aft else None,
            aft_base.get("share") if aft_base else None,
            aft_share - aft_base.get("share", 0) if aft_share is not None and aft_base else None,
        ])
    set_cell(sheet, 23, 1, "本期核心结论")
    for idx, line in enumerate(summary, start=24):
        set_cell(sheet, idx, 1, str(idx - 23))
        set_cell(sheet, idx, 2, line)


def write_stat_sheet(workbook, sheet_name, stage, rows, total, days, buyers):
    sheet = workbook[sheet_name]
    clear_values(sheet)
    set_cell(sheet, 1, 1, f"本期{stage}需求统计")
    set_cell(sheet, 2, 1, "按同一买家 + 阶段 + 需求词条 + 当前周期去重；占比为阶段内占比。")
    write_row(sheet, 4, ["正式需求总数", total, "统计天数", days, "需求覆盖买家", buyers, None, None, None, None, None, None])
    write_row(sheet, 6, ["排名", "需求词条", "同人去重人数", "占比", "日均人数", "原始命中", "重复次数", "优先级", "严格定义", "涉及买家示例", "代表原话", "建议动作"])
    output = []
    for idx, item in enumerate(rows, 1):
        output.append([idx, item["label"], item["count"], item["count"] / total if total else 0, item["count"] / days if days else 0, item["raw_hits"], item["duplicates"], item["priority"], item["definition"], item["buyers"], item["quote"], item["action"]])
    write_rows(sheet, 7, output or [[1, "本期未识别", 0, 0, 0, 0, 0, "—", "—", "—", "—", "—"]], 7)


def validate_stage_ranking(rows, stage):
    expected = sorted(rows, key=lambda item: (-item["count"], item["label"]))
    actual = [(item["label"], item["count"]) for item in rows]
    wanted = [(item["label"], item["count"]) for item in expected]
    if actual != wanted:
        raise RuntimeError(f"{stage}需求排名未按同人去重人数降序排列：{actual}")


def build_compare_rows(rows, total, days, baseline_stage=None, same_period=None):
    baseline_stage = baseline_stage or {"rows": {}, "total": 0, "days": 0}
    baseline_rows = baseline_stage.get("rows", {})
    baseline_days = baseline_stage.get("days", 0) or 0
    if same_period is None:
        same_period = bool(baseline_days) and days == baseline_days
    current = {item["label"]: item for item in rows}
    labels = list(current)
    labels.extend(label for label in baseline_rows if label not in current)
    output = []
    for label in labels:
        item = current.get(label)
        base = baseline_rows.get(label, {})
        current_count = item["count"] if item else 0
        current_share = current_count / total if total else 0
        base_count = base.get("count", 0) or 0
        base_share = base.get("share", 0) or 0
        current_daily = current_count / days if days else 0
        base_daily = base.get("daily", 0) or 0
        share_change = current_share - base_share
        daily_change_rate = current_daily / base_daily - 1 if base_daily else None
        if item and not base:
            judgment = "本期新增词条"
        elif base and not item:
            judgment = "本期未出现"
        elif abs(share_change) >= 0.03:
            judgment = "结构明显上升" if share_change > 0 else "结构明显下降"
        elif not same_period and daily_change_rate is not None and abs(daily_change_rate) >= 0.2:
            judgment = "日均明显上升" if daily_change_rate > 0 else "日均明显下降"
        elif same_period and current_count > base_count:
            judgment = "人数上升"
        elif same_period and current_count < base_count:
            judgment = "人数下降"
        else:
            judgment = "基本持平"
        output.append([label, current_count, current_share, current_daily, base_count, base_share, base_daily, current_count - base_count, share_change, daily_change_rate, judgment, item["action"] if item else base.get("action", "")])
    return output


def write_compare_sheet(workbook, sheet_name, stage, rows, total, days, baseline_stage=None, same_period=None):
    sheet = workbook[sheet_name]
    clear_values(sheet)
    baseline_stage = baseline_stage or {"rows": {}, "total": 0, "days": 0}
    baseline_total = baseline_stage.get("total", 0) or 0
    baseline_days = baseline_stage.get("days", 0) or 0
    set_cell(sheet, 1, 1, f"{stage}需求基准对比")
    set_cell(sheet, 2, 1, "本期沿用上期词条、定义、阶段和去重口径；周期不同时优先比较占比与日均变化率。")
    write_row(sheet, 4, ["本期总需求", total, "本期天数", days, "基准总需求", baseline_total, "基准天数", baseline_days, None, None, None, None])
    write_row(sheet, 6, ["需求词条", "本期人数", "本期占比", "本期日均", "基准人数", "基准占比", "基准日均", "人数变化", "占比变化", "日均变化率", "变化判断", "建议动作"])
    output = build_compare_rows(rows, total, days, baseline_stage, same_period)
    write_rows(sheet, 7, output or [["本期未识别", 0, 0, 0, 0, 0, 0, 0, 0, "V1基准", "—", "—"]], 7)


def write_simple_sheet(workbook, sheet_name, rows):
    sheet = workbook[sheet_name]
    clear_values(sheet)
    write_rows(sheet, 1, rows, 1)


def safe_sheet_title(workbook, value):
    base = re.sub(r"[\\/*?:\[\]]", "_", str(value or "商品"))[:31] or "商品"
    title = base
    suffix = 2
    while title in workbook.sheetnames:
        marker = f"_{suffix}"
        title = base[: 31 - len(marker)] + marker
        suffix += 1
    return title


def cloned_sheet(workbook, source_name, title):
    sheet = workbook.copy_worksheet(workbook[source_name])
    sheet.title = safe_sheet_title(workbook, title)
    clear_values(sheet)
    return sheet


def write_multi_product_overview(workbook, results):
    count_names = {2: "双", 3: "三", 4: "四", 5: "五", 6: "六"}
    count_label = count_names.get(len(results), str(len(results)))
    sheet = cloned_sheet(workbook, "一级主题", f"{count_label}SKU概览")
    set_cell(sheet, 1, 1, f"{count_label}SKU结构对比｜每个聊天文件独立对应一个商品")
    write_row(
        sheet,
        3,
        ["商品/SKU", "有效买家消息", "相关买家", "正式需求", "售前", "售后", "售后占比", "风险", "待复核", "核心判断"],
    )
    for row_number, result in enumerate(results, 4):
        item = result["summary"]
        demand_total = item["demand_count"] or 0
        top_presale = next(iter(sorted_stats([row for row in result["stats"] if row["stage"] == "售前"])), None)
        top_aftersale = next(iter(sorted_stats([row for row in result["stats"] if row["stage"] == "售后"])), None)
        judgment = f"售前重点：{top_presale['label'] if top_presale else '暂无'}；售后重点：{top_aftersale['label'] if top_aftersale else '暂无'}。"
        ensure_rows(sheet, row_number, min(2, sheet.max_row))
        write_row(
            sheet,
            row_number,
            [
                item["product"],
                item["valid_messages"],
                item["buyers"],
                demand_total,
                item["presale_count"],
                item["aftersale_count"],
                item["aftersale_count"] / demand_total if demand_total else 0,
                item["risk_count"],
                item["review_count"],
                judgment,
            ],
        )
    return sheet.title


def write_group_board(workbook, results, presale, aftersale):
    sheet = cloned_sheet(workbook, "售前需求统计", "商品组总盘")
    products = [result["summary"]["product"] for result in results]
    product_stats = {
        result["summary"]["product"]: {
            (item["stage"], item["label"]): item["count"]
            for item in result["stats"]
        }
        for result in results
    }
    set_cell(sheet, 1, 1, "商品组总盘｜统一词条体系，分商品展示人数")

    def write_stage(start_row, stage, rows):
        total = sum(item["count"] for item in rows)
        ensure_rows(sheet, start_row, 1)
        ensure_rows(sheet, start_row + 1, 6)
        set_cell(sheet, start_row, 1, f"商品组{stage}需求")
        write_row(
            sheet,
            start_row + 1,
            ["排名", "需求词条", "去重需求数", "占比", "优先级", "严格定义", "建议动作", *products],
        )
        for index, item in enumerate(rows, 1):
            ensure_rows(sheet, start_row + 1 + index, 7)
            write_row(
                sheet,
                start_row + 1 + index,
                [
                    index,
                    item["label"],
                    item["count"],
                    item["count"] / total if total else 0,
                    item["priority"],
                    item["definition"],
                    item["action"],
                    *[product_stats[product].get((stage, item["label"]), 0) for product in products],
                ],
            )
        return start_row + len(rows) + 3

    aftersale_start = write_stage(3, "售前", presale)
    write_stage(aftersale_start, "售后", aftersale)
    return sheet.title


def baseline_for_product(baseline_data, product, result_count):
    if not baseline_data:
        return None
    products = baseline_data.get("products") or {}
    if not products and result_count == 1 and baseline_data.get("taxonomy"):
        return {"售前": baseline_data["售前"], "售后": baseline_data["售后"]}
    normalized = re.sub(r"\W+", "", product).upper()
    tokens = set(re.findall(r"(?i)(?:\d+[A-Z][A-Z0-9]*|[A-Z]+\d+[A-Z0-9]*)", product))
    best = None
    best_score = 0
    for baseline_name, baseline_product in products.items():
        baseline_normalized = re.sub(r"\W+", "", baseline_name).upper()
        baseline_tokens = set(re.findall(r"(?i)(?:\d+[A-Z][A-Z0-9]*|[A-Z]+\d+[A-Z0-9]*)", baseline_name))
        score = 3 if normalized == baseline_normalized else 2 if tokens & baseline_tokens else 1 if normalized in baseline_normalized or baseline_normalized in normalized else 0
        if score > best_score:
            best = baseline_product
            best_score = score
    return best


def write_product_sheet(workbook, result, baseline_product=None, same_period=None):
    item = result["summary"]
    product = item["product"]
    days = len({message.get("date") for message in result["valid_messages"] if message.get("date")}) or 1
    sheet = cloned_sheet(workbook, "售前需求统计", product)
    set_cell(sheet, 1, 1, f"{product}｜独立客服需求分析")
    set_cell(
        sheet,
        2,
        1,
        f"正式需求 {item['demand_count']}；需求买家 {item['buyers']}；售前 {item['presale_count']}；售后 {item['aftersale_count']}；真实质量风险 {item['risk_count']}。",
    )

    def write_stage(start_row, stage):
        rows = sorted_stats([row for row in result["stats"] if row["stage"] == stage])
        total = sum(row["count"] for row in rows)
        ensure_rows(sheet, start_row, 1)
        ensure_rows(sheet, start_row + 1, 6)
        set_cell(sheet, start_row, 1, f"{stage}V1")
        write_row(
            sheet,
            start_row + 1,
            ["排名", "需求词条", "去重人数", "阶段内占比", "日均", "原始命中", "重复次数", "优先级", "严格定义", "涉及买家", "代表原话", "建议动作"],
        )
        for index, row in enumerate(rows, 1):
            ensure_rows(sheet, start_row + 1 + index, 7)
            write_row(
                sheet,
                start_row + 1 + index,
                [index, row["label"], row["count"], row["count"] / total if total else 0, row["count"] / days, row["raw_hits"], row["duplicates"], row["priority"], row["definition"], row["buyers"], row["quote"], row["action"]],
            )
        return start_row + len(rows) + 3

    aftersale_start = write_stage(4, "售前")
    next_row = write_stage(aftersale_start, "售后")
    if baseline_product:
        for stage in ("售前", "售后"):
            rows = sorted_stats([row for row in result["stats"] if row["stage"] == stage])
            total = sum(row["count"] for row in rows)
            baseline_stage = baseline_product.get(stage) or {"rows": {}, "total": 0, "days": 0}
            comparison = build_compare_rows(rows, total, days, baseline_stage, same_period)
            ensure_rows(sheet, next_row, 1)
            ensure_rows(sheet, next_row + 1, 6)
            set_cell(sheet, next_row, 1, f"{stage}基准对比")
            write_row(sheet, next_row + 1, ["需求词条", "本期人数", "本期占比", "本期日均", "基准人数", "基准占比", "基准日均", "人数变化", "占比变化", "日均变化率", "变化判断", "建议动作"])
            for offset, row in enumerate(comparison, 1):
                ensure_rows(sheet, next_row + 1 + offset, 7)
                write_row(sheet, next_row + 1 + offset, row)
            next_row += len(comparison) + 3
    return sheet.title


def write_daily_trend(workbook, results):
    sheet = cloned_sheet(workbook, "一级主题", "每日趋势")
    products = [result["summary"]["product"] for result in results]
    dates = sorted(
        {
            item.get("date")
            for result in results
            for item in result["detail"]
            if item.get("date")
        }
    )
    counts = Counter(
        (item["product"], item.get("date"))
        for result in results
        for item in result["detail"]
        if item.get("date")
    )
    set_cell(sheet, 1, 1, "每日首次需求趋势｜按商品拆分")
    write_row(sheet, 3, ["日期", *products, "总需求"])
    for row_number, date in enumerate(dates, 4):
        product_counts = [counts[(product, date)] for product in products]
        ensure_rows(sheet, row_number, min(2, sheet.max_row))
        write_row(sheet, row_number, [date, *product_counts, sum(product_counts)])
    return sheet.title


def write_multi_product_sheets(workbook, results, presale, aftersale, baseline_data=None, same_period=None):
    if len(results) <= 1:
        return []
    created = [
        write_multi_product_overview(workbook, results),
        write_group_board(workbook, results, presale, aftersale),
    ]
    created.extend(
        write_product_sheet(
            workbook,
            result,
            baseline_for_product(baseline_data, result["summary"]["product"], len(results)),
            same_period,
        )
        for result in results
    )
    created.append(write_daily_trend(workbook, results))
    return created


def write_scope_sheet(workbook, data, results, baseline_data):
    sheet = cloned_sheet(workbook, "一级主题", "口径说明")
    products = [result["summary"]["product"] for result in results]
    valid_messages = [message for result in results for message in result["valid_messages"]]
    dates = sorted({message.get("date") for message in valid_messages if message.get("date")})
    buyers = {message["sender"] for message in valid_messages}
    baseline_period = "无历史基准"
    if baseline_data.get("start_date") and baseline_data.get("end_date"):
        baseline_period = f"{baseline_data['start_date']} 至 {baseline_data['end_date']}"
    rows = [
        ["口径项", "本次规则", "说明"],
        ["分析对象", data.get("product_name") or "客服聊天需求分析", "、".join(products)],
        ["统计时间", f"{dates[0]} 至 {dates[-1]}" if dates else "未识别", f"按聊天实际日期，共 {len(dates)} 个有消息日期"],
        ["统计对象", f"{len(valid_messages)} 条有效买家消息 / {len(buyers)} 名买家", "排除客服坐席、系统卡片、订单卡片、链接和纯噪声"],
        ["售前定义", "购买决策前", "参数、规格、款式、价格、政策、购买前物流和安全顾虑"],
        ["售后定义", "已购/收货/实操/故障", "使用教程、安装、实际效果、故障、退换货和履约处理"],
        ["去重规则", "商品 + 买家 + 阶段 + 需求词条 + 当前周期", "同一商品内同一买家重复问同一词条只计 1 次，原话进入重复记录"],
        ["多需求拆分", "允许", "同一消息包含多个独立诉求时分别计入对应词条"],
        ["质量风险", "必须本人实际发生并完成二审", "购买前担忧、客服话术和正常现象不进入正式风险"],
        ["上期基准", baseline_period, "有基准时沿用词条、定义、阶段和去重规则，只新增旧词条无法覆盖的需求"],
        ["完整性门禁", f"{data.get('llm_usage', {}).get('analyzed_messages', 0)}/{data.get('llm_usage', {}).get('expected_messages', 0)}", "未完成全部有效消息时禁止发布正式 Excel"],
    ]
    if str((data.get("manual_baseline") or {}).get("text") or "").strip():
        rows.append(
            [
                "上期基准来源",
                "手工录入",
                "仅统计录入的词条与数字；未录入项按无基准处理，定义缺失的词条按本期聊天重新校准",
            ]
        )
    write_rows(sheet, 1, rows, 1)
    return sheet.title


def write_valid_messages_sheet(workbook, results):
    sheet = cloned_sheet(workbook, "原始数据", "有效买家消息")
    rows = [["商品", "买家", "日期", "时间", "会话序号", "源行", "有效买家原话"]]
    rows.extend(
        [message["product"], message["sender"], message.get("date", ""), message.get("time", ""), message.get("conversation_id", 0), message.get("source_line", ""), message["text"]]
        for result in results
        for message in result["valid_messages"]
    )
    write_rows(sheet, 1, rows, 1)
    if sheet.max_row > len(rows):
        sheet.delete_rows(len(rows) + 1, sheet.max_row - len(rows))
    return sheet.title


def build_summary(results):
    all_summary = [item["summary"] for item in results]
    total_demand = sum(item["demand_count"] for item in all_summary)
    total_presale = sum(item["presale_count"] for item in all_summary)
    total_aftersale = sum(item["aftersale_count"] for item in all_summary)
    total_risks = sum(item["risk_count"] for item in all_summary)
    presale = aggregate_stats(results, "售前")
    aftersale = aggregate_stats(results, "售后")
    top_pre = presale[0]["label"] if presale else "暂无"
    top_after = aftersale[0]["label"] if aftersale else "暂无"
    return [
        f"本期共识别正式需求 {total_demand} 条，其中售前 {total_presale} 条、售后 {total_aftersale} 条。",
        f"售前最高频需求是“{top_pre}”，建议优先优化详情页、型号/功能说明和客服自动回复。",
        f"售后最高频需求是“{top_after}”，说明用户收到货后的使用学习或售后承接压力需要优先处理。",
        f"本期识别安全/质量风险候选 {total_risks} 条，已进入《安全质量风险》Sheet，需人工二次复核。",
        "建议先处理 P0 词条：使用方法、故障/效果、退款退货和安全风险。"
    ]


def write_workbook(data, output_path):
    baseline_files = data.get("baseline_files") or []
    uploaded_baseline = next(
        (
            Path(item.get("path", ""))
            for item in baseline_files
            if Path(item.get("path", "")).suffix.lower() in (".xlsx", ".xls") and Path(item.get("path", "")).exists()
        ),
        None,
    )
    # Baselines provide historical data only. Report formatting always comes from
    # the single formal mother workbook so old product layouts cannot leak in.
    template_path = TEMPLATE_XLSX
    if not template_path.exists():
        raise FileNotFoundError(f"找不到正式报告母版：{template_path}")
    manual_baseline = data.get("manual_baseline") or {}
    manual_baseline_text = str(manual_baseline.get("text") or "").strip()
    baseline_data = read_baseline_data(None)
    if uploaded_baseline:
        baseline_data = read_baseline_data(uploaded_baseline)
        if not baseline_data.get("taxonomy"):
            raise RuntimeError("上期基准未读取到词条、定义和阶段，完整性门禁已阻止无口径对比。")
    elif manual_baseline_text:
        baseline_data = parse_manual_baseline(
            manual_baseline_text,
            days=manual_baseline.get("days"),
            period=manual_baseline.get("period"),
        )
    baseline_metrics = None
    if uploaded_baseline or manual_baseline_text:
        baseline_metrics = {
            "period": {
                "start_date": baseline_data.get("start_date", ""),
                "end_date": baseline_data.get("end_date", ""),
            }
        }
        for stage in ("售前", "售后"):
            baseline_metrics[stage] = {
                "total": baseline_data[stage]["total"],
                "days": baseline_data[stage]["days"],
                "rows": [
                    {
                        "label": label,
                        "count": item["count"],
                        "share": item["share"],
                        "daily": item["daily"],
                    }
                    for label, item in sorted(
                        baseline_data[stage]["rows"].items(),
                        key=lambda pair: (-pair[1]["count"], pair[0]),
                    )
                ],
            }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    staging_path = output_path.with_name(f"{output_path.stem}__staging.xlsx")
    shutil.copy2(template_path, staging_path)
    workbook = load_workbook(staging_path)
    ensure_sheets(workbook)

    chat_files = data.get("chat_files") or data.get("saved_files") or []
    if is_generic_product_name(data.get("product_name")) and uploaded_baseline:
        inferred_product = product_from_baseline_file(uploaded_baseline)
        if inferred_product:
            data["product_name"] = inferred_product
    entries = []
    for item in chat_files:
        file_path = Path(item.get("path", ""))
        if not file_path.exists():
            continue
        product = product_from_file(file_path)
        if len(chat_files) == 1 and not is_generic_product_name(data.get("product_name")):
            product = data["product_name"]
        raw = parse_messages(product, file_path)
        entries.append({"product": product, "raw": raw, "valid": valid_customer_messages(raw)})
    if uploaded_baseline and baseline_data.get("products"):
        baseline_products = baseline_data["products"]
        if len(entries) != len(baseline_products):
            raise RuntimeError(
                "多商品基准完整性校验失败：本次上传 "
                f"{len(entries)} 个聊天记录，上期基准包含 {len(baseline_products)} 个商品。"
            )
        matched = [baseline_for_product(baseline_data, entry["product"], len(entries)) for entry in entries]
        missing = [entry["product"] for entry, item in zip(entries, matched) if item is None]
        if missing:
            raise RuntimeError(
                "多商品基准完整性校验失败：以下商品无法匹配上期基准：" + "、".join(missing)
            )
        if len({id(item) for item in matched}) != len(matched):
            raise RuntimeError("多商品基准完整性校验失败：多个聊天文件匹配到了同一个上期商品。")
    if entries:
        results, taxonomy, summary, llm_usage = analyze_entries_with_llm(
            entries,
            task_context=data.get("user_message") or data.get("product_name") or "客服聊天需求分析",
            baseline_context=baseline_data.get("context", ""),
            reference_taxonomy=baseline_data.get("taxonomy") or None,
            baseline_metrics=baseline_metrics,
        )
        data["analysis_engine"] = "llm_api"
        data["taxonomy"] = taxonomy
        data["llm_usage"] = llm_usage
        data["products"] = [result["summary"]["product"] for result in results]
    else:
        results = []
        summary = ["未找到可解析的聊天记录文件。"]
    days = len(set(msg["date"] for result in results for msg in result["valid_messages"] if msg.get("date"))) or 1
    buyers = len({message["sender"] for result in results for message in result["valid_messages"]})
    presale = aggregate_stats(results, "售前") if results else []
    aftersale = aggregate_stats(results, "售后") if results else []
    total_presale = sum(item["count"] for item in presale)
    total_aftersale = sum(item["count"] for item in aftersale)
    validate_stage_ranking(presale, "售前")
    validate_stage_ranking(aftersale, "售后")
    current_dates = sorted(
        {
            message.get("date")
            for result in results
            for message in result["valid_messages"]
            if message.get("date")
        }
    )
    same_period = bool(
        uploaded_baseline
        and current_dates
        and baseline_data.get("start_date") == current_dates[0]
        and baseline_data.get("end_date") == current_dates[-1]
    )

    data["summary"] = summary
    write_overview(workbook, data, results, summary, baseline_data)
    write_simple_sheet(workbook, "一级主题", aggregate_theme_rows(results))
    write_stat_sheet(workbook, "售前需求统计", "售前", presale, total_presale, days, buyers)
    write_stat_sheet(workbook, "售后需求统计", "售后", aftersale, total_aftersale, days, buyers)
    write_compare_sheet(workbook, "售前基准对比", "售前", presale, total_presale, days, baseline_data["售前"], same_period)
    write_compare_sheet(workbook, "售后基准对比", "售后", aftersale, total_aftersale, days, baseline_data["售后"], same_period)
    data["multi_product_sheets"] = write_multi_product_sheets(workbook, results, presale, aftersale, baseline_data, same_period)
    data["supplemental_sheets"] = [
        write_scope_sheet(workbook, data, results, baseline_data),
        write_valid_messages_sheet(workbook, results),
    ]

    baseline_labels = {item["label"] for item in baseline_data.get("taxonomy") or []}
    new_items = [
        item
        for item in presale + aftersale
        if not uploaded_baseline or item["label"] not in baseline_labels
    ]
    write_simple_sheet(
        workbook,
        "新增词条",
        [["阶段", "新增词条", "本期人数", "本期占比", "优先级", "严格定义", "建议动作"]]
        + [
            [
                item["stage"],
                item["label"],
                item["count"],
                item["count"] / max(1, total_presale if item["stage"] == "售前" else total_aftersale),
                item["priority"],
                item["definition"],
                item["action"],
            ]
            for item in new_items
        ],
    )
    stat_meta = {(item["product"], item["label"]): item for result in results for item in result["stats"]}
    p0_items = [
        detail
        for result in results
        for detail in result["detail"]
        if stat_meta.get((detail["product"], detail["label"]), {}).get("priority") == "P0"
    ]
    write_simple_sheet(workbook, "P0风险明细", [["ID", "商品", "需求词条", "买家", "日期", "时间", "代表原话", "全部证据", "源行", "优先级", "建议动作"]] + [[idx, item["product"], item["label"], item["sender"], item["date"], item["time"], item["text"], item["all_text"], item["source_line"], stat_meta[(item["product"], item["label"])]["priority"], stat_meta[(item["product"], item["label"])]["action"]] for idx, item in enumerate(p0_items, 1)])
    write_simple_sheet(workbook, "安全质量风险", [["风险ID", "商品", "风险词条", "买家", "日期", "时间", "代表原话", "全部证据", "源行", "原始命中", "建议动作"]] + [[idx, item["product"], item["risk_label"], item["sender"], item["date"], item["time"], item["text"], item.get("all_text", item["text"]), item["source_line"], item.get("raw_hit", 1), item["action"]] for idx, item in enumerate([risk for r in results for risk in r["risks"]], 1)])
    write_simple_sheet(workbook, "逐条需求明细", [["明细ID", "商品", "阶段", "需求词条", "买家", "日期", "时间", "代表原话", "同需求全部原话", "源行", "原始命中", "一级主题", "分析来源"]] + [[idx, item["product"], item["stage"], item["label"], item["sender"], item["date"], item["time"], item["text"], item["all_text"], item["source_line"], item["raw_hit"], item["theme"], "AI语义分析"] for idx, item in enumerate([d for r in results for d in r["detail"]], 1)])
    write_simple_sheet(workbook, "同人重复记录", [["商品", "阶段", "需求词条", "买家", "日期", "时间", "重复原话", "重复源行", "保留原话", "剔除原因"]] + [[item["product"], item["stage"], item["label"], item["sender"], item["date"], item["time"], item["text"], item["source_line"], item["kept_text"], "同一商品、同一买家、同一阶段、同一需求词条，本周期只计1次"] for item in [dup for r in results for dup in r["duplicates"]]])
    write_simple_sheet(workbook, "待人工复核", [["商品", "买家", "日期", "时间", "待复核原话", "源行", "复核原因", "状态"]] + [[item["product"], item["sender"], item["date"], item["time"], item["text"], item["source_line"], item["reason"], "待人工确认"] for item in [rev for r in results for rev in r["reviews"]]])
    current_counts = {
        "售前": {item["label"]: item["count"] for item in presale},
        "售后": {item["label"]: item["count"] for item in aftersale},
    }
    taxonomy_rows = []
    for item in taxonomy:
        stages = item.get("stages") or [item.get("stage", "售后")]
        taxonomy_rows.append(
            [
                item["label"],
                "、".join(stages),
                item.get("theme") or infer_theme_from_label(item["label"]),
                item["definition"],
                item["priority"],
                current_counts["售前"].get(item["label"], 0),
                current_counts["售后"].get(item["label"], 0),
                "沿用上期V1" if item["label"] in baseline_labels else "AI本期新增",
                item["action"],
            ]
        )
    write_simple_sheet(workbook, "词条与口径", [["需求词条", "本期阶段", "一级主题", "严格定义", "优先级", "售前人数", "售后人数", "生成状态", "建议动作"]] + taxonomy_rows)
    if uploaded_baseline:
        baseline_output = []
        for stage in ("售前", "售后"):
            stage_rows = sorted(
                baseline_data[stage]["rows"].items(),
                key=lambda pair: (-pair[1]["count"], pair[0]),
            )
            for rank, (label, item) in enumerate(stage_rows, 1):
                baseline_output.append([stage, rank, label, item["count"], item["share"], item["daily"], item["priority"], item.get("definition", ""), item.get("action", "")])
    else:
        baseline_output = []
        for stage, stage_rows, stage_total in (("售前", presale, total_presale), ("售后", aftersale, total_aftersale)):
            for rank, item in enumerate(stage_rows, 1):
                baseline_output.append([stage, rank, item["label"], item["count"], item["count"] / max(1, stage_total), item["count"] / days if days else 0, item["priority"], item["definition"], item["action"]])
    write_simple_sheet(workbook, "上期基准", [["阶段", "排名", "需求词条", "基准人数", "基准占比", "基准日均", "优先级", "严格定义", "建议动作"]] + baseline_output)
    write_simple_sheet(workbook, "原始数据", [["产品", "会话序号", "发送方", "是否买家有效消息", "日期", "时间", "源行", "原文"]] + [[msg["product"], msg["conversation_id"], msg["sender"], "是" if is_valid_customer_message(msg) else "否", msg["date"], msg["time"], msg["source_line"], msg["text"]] for msg in [m for entry in entries for m in entry["raw"]]])

    workbook.save(staging_path)
    write_values_into_exact_template(template_path, staging_path, output_path)
    staging_path.unlink(missing_ok=True)
    return summary


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: generate_report.py report_data.json output.xlsx")
    data_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    data = json.loads(data_path.read_text(encoding="utf-8"))
    summary = write_workbook(data, output_path)
    data["summary"] = summary
    data_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
