import json
import sys
from pathlib import Path

from openpyxl import load_workbook


PRODUCT_ROWS = {
    "856": 10,
    "366": 11,
    "五合一": 12,
}


def nonempty_rows(ws, start_row=1):
    return sum(
        1
        for row in ws.iter_rows(min_row=start_row, values_only=True)
        if any(value is not None and str(value).strip() for value in row)
    )


def workbook_metrics(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        overview = workbook["分析总览"]
        products = {}
        for product, row in PRODUCT_ROWS.items():
            products[product] = {
                "buyers": int(overview.cell(row, 2).value or 0),
                "demand": int(overview.cell(row, 3).value or 0),
                "presale": int(overview.cell(row, 4).value or 0),
                "aftersale": int(overview.cell(row, 5).value or 0),
                "risk": int(overview.cell(7, {"856": 4, "366": 8, "五合一": 12}[product]).value or 0),
                "top_presale": str(overview.cell(row, 10).value or "").rsplit(" ", 1)[0],
                "top_aftersale": str(overview.cell(row, 11).value or "").rsplit(" ", 1)[0],
            }
        dictionary = workbook["V1词条字典"]
        labels = {
            str(row[0]).strip()
            for row in dictionary.iter_rows(min_row=2, values_only=True)
            if row and row[0]
        }
        values = [
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        return {
            "products": products,
            "valid_messages": nonempty_rows(workbook["有效买家消息"], 2),
            "detail_rows": nonempty_rows(workbook["逐条需求明细"], 2),
            "risk_rows": nonempty_rows(workbook["质量安全风险"], 5),
            "review_rows": nonempty_rows(workbook["待人工复核"], 2),
            "taxonomy_count": len(labels),
            "labels": sorted(labels),
            "placeholder_hits": sum(
                1
                for value in values
                if "待真实分析写入" in value or "模型 API 未完成自动分类" in value
            ),
        }
    finally:
        workbook.close()


def ratio(candidate, baseline):
    if not baseline:
        return 1.0 if not candidate else float("inf")
    return candidate / baseline


def compare(baseline, candidate):
    failures = []
    checks = {}
    valid_ratio = ratio(candidate["valid_messages"], baseline["valid_messages"])
    checks["valid_message_ratio"] = valid_ratio
    if candidate["valid_messages"] != baseline["valid_messages"]:
        failures.append("有效买家消息未与基准完全一致")

    total_baseline = sum(item["demand"] for item in baseline["products"].values())
    total_candidate = sum(item["demand"] for item in candidate["products"].values())
    demand_ratio = ratio(total_candidate, total_baseline)
    checks["total_demand_ratio"] = demand_ratio
    if not 0.9 <= demand_ratio <= 1.1:
        failures.append("正式需求总数与基准偏差超过 10%")

    for product in PRODUCT_ROWS:
        for metric in ("presale", "aftersale"):
            value = ratio(candidate["products"][product][metric], baseline["products"][product][metric])
            checks[f"{product}_{metric}_ratio"] = value
            if not 0.8 <= value <= 1.2:
                failures.append(f"{product} {metric} 与基准偏差超过 20%")
        for metric in ("top_presale", "top_aftersale"):
            if candidate["products"][product][metric] != baseline["products"][product][metric]:
                failures.append(f"{product} {metric} 未与基准一致")

    risk_ratio = ratio(candidate["risk_rows"], baseline["risk_rows"])
    checks["risk_ratio"] = risk_ratio
    if not 0.85 <= risk_ratio <= 1.15:
        failures.append("强风险数量与基准偏差超过 15%")

    label_coverage = len(set(candidate["labels"]) & set(baseline["labels"])) / max(1, len(baseline["labels"]))
    checks["taxonomy_coverage"] = label_coverage
    if label_coverage < 1.0:
        failures.append("正式 V1 词条覆盖率未达到 100%")
    if candidate["placeholder_hits"]:
        failures.append("报告仍包含接口失败或待写入占位内容")

    return {"passed": not failures, "checks": checks, "failures": failures}


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: qa_compare_curling_baseline.py baseline.xlsx candidate.xlsx")
    baseline_path = Path(sys.argv[1]).resolve()
    candidate_path = Path(sys.argv[2]).resolve()
    baseline = workbook_metrics(baseline_path)
    candidate = workbook_metrics(candidate_path)
    result = {
        "baseline": str(baseline_path),
        "candidate": str(candidate_path),
        "baseline_metrics": baseline,
        "candidate_metrics": candidate,
        **compare(baseline, candidate),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
