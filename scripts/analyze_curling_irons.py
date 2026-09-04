import json
import os
import re
import sys
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from exact_excel_writer import write_values_into_exact_template
from llm_analysis import analyze_entries_with_llm


INPUTS = [
    ("五合一卷发棒", Path(r"E:/新2/赤兔历史分析结果/五合一聊天记录.log")),
    ("366卷发棒", Path(r"E:/新2/赤兔历史分析结果/366聊天记录.log")),
    ("856卷发棒", Path(r"E:/新2/赤兔历史分析结果/856聊天记录.log")),
]

OUTPUT_ROOT = Path(os.environ.get("CHITU_OUTPUT_ROOT", r"E:/新2/赤兔历史分析结果/03_最终报告/2026/08"))
TEMPLATE_XLSX = Path(
    os.environ.get(
        "CHITU_CURLING_TEMPLATE_XLSX",
        r"E:/新2/赤兔聊天分析工作台/templates/excel/三SKU首次基准_卷发棒标准模板.xlsx",
    )
)
TAXONOMY_XLSX = Path(os.environ.get("CHITU_CURLING_TAXONOMY_XLSX", str(TEMPLATE_XLSX)))


def resolve_inputs():
    spec = os.environ.get("CHITU_CURLING_INPUTS_JSON")
    if not spec:
        return INPUTS
    items = json.loads(spec)
    resolved = []
    for item in items:
        product = item.get("product")
        file_path = item.get("path")
        if not product or not file_path:
            raise ValueError("CHITU_CURLING_INPUTS_JSON must contain product and path")
        resolved.append((product, Path(file_path)))
    return resolved

RED = "B5322D"
DARK_RED = "8D241F"
LIGHT_RED = "FFF5F2"
LIGHT_BLUE = "F2F7FF"
TEXT = "1F2633"
BORDER = "D9DEE7"

thin_border = Border(
    left=Side(style="thin", color=BORDER),
    right=Side(style="thin", color=BORDER),
    top=Side(style="thin", color=BORDER),
    bottom=Side(style="thin", color=BORDER),
)

HEADER_RE = re.compile(r"^(.+?)\s+(\d{4}-\d{2}-\d{2})\s+(\d{2}:\d{2}:\d{2})\s*$")
URL_RE = re.compile(r"https?://\S+")
SERVICE_SENDER_PATTERNS = (
    re.compile(r"^丹橘个护健康沭阳\d+$", re.I),
    re.compile(r"^jimi_vender_\d+$", re.I),
)

INVALID_EXACT = {
    "你好",
    "您好",
    "在吗",
    "好的",
    "好",
    "嗯",
    "谢谢",
    "谢谢你",
    "OK",
    "ok",
    "好吧",
    "哦哦",
    "是的",
    "好了",
    "可以",
    "可以了",
    "谢谢啦",
    "好的好的",
    "嗯嗯",
    "谢谢！",
    "谢谢，",
    "谢谢亲",
    "谢谢您",
    "谢谢哦",
    "谢谢美女",
    "谢谢你了",
    "谢谢知道啦",
    "谢谢，辛苦了",
    "知道了",
    "同意",
    "评价了",
    "已评价",
    "好评了",
    "人工",
    "人工客服",
    "转人工",
    "用户发起转人工",
    "客服",
    "在不",
    "不用了吧",
    "嗯没事",
    "哦，",
    "额！",
    "啊？",
    "^_^n",
    "👌🏻",
    "谢谢 比较急",
    "谢谢亲，款已退回",
    "谢谢，小哥离开了",
    "#E-s25那我还得寄回去",
}

INVALID_TEXT_PATTERNS = [
    re.compile(r"^【此消息为欢迎卡片或富文本模板答案，聊天记录中暂不支持展示】$"),
    re.compile(r"^(?:查询|咨询)?订单号[：:].*?(?:商品(?:号|编号|id)|商品ID)[：:]", re.I),
    re.compile(r"^您选择了一个订单进行咨询（订单号：.*?；商品编号：.*?）$"),
    re.compile(r"^(?:#E-[A-Za-z0-9]+)+$"),
    re.compile(r"^[？?!！。,.，…~～^_\-\s]+$"),
]

LABEL_RULES = [
    ("退款退货", "售后", "售后服务", ["退款", "退货", "退了", "退掉", "退换", "申请退", "拦截", "拒收", "运费险"]),
    ("拆卸/更换配件", "售后", "使用学习", ["拆卸", "拆下", "拆装", "怎么拆", "更换", "换头", "替换", "安装", "卡扣"]),
    ("使用方法/教程", "售后", "使用学习", ["怎么用", "如何使用", "使用方法", "教程", "视频", "操作", "怎么开", "怎么弄", "怎么卷", "怎么夹", "开机", "按键"]),
    ("不加热/温度低", "售后", "加热温控", ["不热", "没温度", "温度低", "不加热", "加热慢", "温度上不去", "200度", "几乎没什么温度"]),
    ("不出雾/喷雾异常", "售后", "喷雾系统", ["不出雾", "没有雾", "没雾", "喷雾", "出雾", "水箱", "加水", "漏水"]),
    ("造型效果不好", "售后", "造型效果", ["没效果", "效果不好", "卷不出来", "卷不上", "不定型", "不持久", "夹不直", "拉不直", "不好用"]),
    ("故障/无法工作", "售后", "质量故障", ["坏了", "故障", "无法使用", "没反应", "失灵", "断电", "插电", "电源", "指示灯", "报警"]),
    ("好评/返现/补偿", "售后", "售后服务", ["好评", "返现", "返钱", "返款", "奖励金", "打款", "补偿", "红包"]),
    ("缺件/赠品/发票", "售后", "售后服务", ["缺", "少了", "赠品", "漏发", "发票", "补发"]),
    ("卷发/直发功能", "售前", "功能咨询", ["能做卷发", "可以卷发", "直卷", "直发", "卷发", "刘海", "内扣", "蓬松", "造型"]),
    ("款式/型号区别", "售前", "购买决策", ["区别", "哪款", "型号", "几款", "套餐", "五合一", "三合一", "366", "856", "选哪个"]),
    ("温度/档位/预热", "售前", "参数规格", ["多少度", "温度", "档位", "预热", "加热", "恒温"]),
    ("是否伤发/发质适用", "售前", "安全顾虑", ["伤头发", "伤发", "适合", "发质", "细软", "粗硬", "会不会伤", "烫伤"]),
    ("发货时效/物流", "售前", "交易政策", ["发货", "送到", "多久到", "物流", "次日达", "地址", "什么时候到"]),
    ("价格优惠/活动", "售前", "交易政策", ["优惠", "券", "活动", "便宜", "价格", "保价", "多少钱"]),
    ("退换/试用政策", "售前", "交易政策", ["运费险", "试用", "七天", "7天", "无理由", "质保", "保修", "售后"]),
    ("参数/尺寸/重量", "售前", "参数规格", ["重量", "尺寸", "多大", "长度", "功率", "净重", "规格"]),
    ("配件/赠品/耗材", "售前", "配件套餐", ["配件", "赠品", "梳子", "夹子", "收纳", "水", "精油"]),
]

RISK_RULES = [
    ("烫伤/安全伤害", ["烫伤", "烫到", "烫手", "烫脸", "烫头皮", "烧焦", "糊味", "冒烟"]),
    ("使用后伤发", ["头发焦", "烧头发", "伤头发", "头发坏", "断发", "掉发"]),
    ("疑似漏电/电流", ["漏电", "触电", "电流", "麻手"]),
    ("异常发热/冒烟", ["异常发热", "很烫", "冒烟", "烧了"]),
]

ACTION_MAP = {
    "使用方法/教程": "下单后和售后首触达自动发送 30 秒使用视频，包装内放步骤卡。",
    "拆卸/更换配件": "把换头/拆卸做成单独短视频，客服自动回复优先推送。",
    "不加热/温度低": "客服先排查预热时间、档位、使用方式；集中样本需品质复测温控。",
    "不出雾/喷雾异常": "补充加水、水箱、喷雾开关和清洁说明，异常样本进入售后复核。",
    "造型效果不好": "按发质、发量、温度、停留时间给标准手法，详情页明确效果边界。",
    "退款退货": "继续拆退款原因，区分不会用、效果差、故障和冲动下单，优先干预可挽回退货。",
    "卷发/直发功能": "详情页首屏增加直卷两用/适合造型示例和视频对照。",
    "款式/型号区别": "做型号对比表，明确适用人群、功能、配件和价格差异。",
    "温度/档位/预热": "详情页和说明书统一温度档位、预热时间和适用发质。",
    "是否伤发/发质适用": "明确护发卖点、适用发质和防烫注意事项，降低售前安全顾虑。",
    "发货时效/物流": "自动回复展示仓配、预计送达和偏远地区说明。",
    "价格优惠/活动": "统一优惠入口和活动口径，减少客服重复解释。",
}

PRIORITY_MAP = {
    "退款退货": "P0",
    "不加热/温度低": "P0",
    "不出雾/喷雾异常": "P0",
    "造型效果不好": "P0",
    "使用方法/教程": "P0",
    "拆卸/更换配件": "P0",
    "故障/无法工作": "P0",
    "是否伤发/发质适用": "P1",
    "款式/型号区别": "P1",
}


def normalize(text):
    return re.sub(r"\s+", " ", text or "").strip()


def is_pure_url(text):
    compact = re.sub(r"\s+", "", text or "")
    return bool(URL_RE.fullmatch(compact))


def is_service(sender, text=""):
    sender_normalized = re.sub(r"\s+", "", sender or "")
    return any(pattern.fullmatch(sender_normalized) for pattern in SERVICE_SENDER_PATTERNS)


def is_invalid_customer_text(text):
    if text in INVALID_EXACT:
        return True
    return any(pattern.search(text) for pattern in INVALID_TEXT_PATTERNS)


def parse_log(product, file):
    lines = file.read_text(encoding="utf-8", errors="replace").splitlines()
    messages = []
    conversation_id = 0
    current = None
    source_line = 0
    for raw in lines:
        source_line += 1
        line = raw.strip()
        if not line:
            continue
        if "以下为一通会话" in line:
            conversation_id += 1
            current = None
            continue
        if "会话结束" in line:
            current = None
            continue
        match = HEADER_RE.match(line)
        if match:
            sender, date, time = match.groups()
            current = {
                "product": product,
                "conversation_id": conversation_id,
                "sender": sender.strip(),
                "date": date,
                "time": time,
                "text": "",
                "source_line": source_line,
            }
            messages.append(current)
        elif current:
            if current["text"]:
                current["text"] += "\n"
            current["text"] += line
    return messages


def valid_customer_messages(messages):
    result = []
    for msg in messages:
        # Normalize only for filtering; preserve the parsed buyer text for prompts
        # and Excel so quoted replies, line breaks, and spacing remain intact.
        raw_text = (msg["text"] or "").strip()
        text = normalize(raw_text)
        if not text:
            continue
        if is_pure_url(raw_text):
            continue
        if len(text) <= 1 or is_invalid_customer_text(text):
            continue
        if is_service(msg["sender"], text):
            continue
        copy = dict(msg)
        copy["text"] = raw_text
        result.append(copy)
    return result


def classify_message(text):
    labels = []
    for label, stage, theme, keywords in LABEL_RULES:
        if any(keyword in text for keyword in keywords):
            labels.append((stage, theme, label))
    if not labels:
        return []
    # Deduplicate while preserving order.
    seen = set()
    unique = []
    for item in labels:
        key = (item[0], item[2])
        if key not in seen:
            seen.add(key)
            unique.append(item)
    return unique


def risk_items(text):
    return [(label, keywords) for label, keywords in RISK_RULES if any(keyword in text for keyword in keywords)]


def analyze_product(product, file):
    raw = parse_log(product, file)
    valid = valid_customer_messages(raw)
    detail = []
    duplicates = []
    retained = {}
    raw_hits = Counter()
    buyers_by_label = defaultdict(set)
    evidence_by_label = defaultdict(list)
    risk_rows = []
    review_rows = []

    for msg in valid:
        labels = classify_message(msg["text"])
        if not labels:
            if any(token in msg["text"] for token in ("不", "没", "退", "坏", "怎么", "为什么")):
                review_rows.append(
                    {
                        **msg,
                        "reason": "存在问题语气但未命中明确词条，需人工判断是否为有效需求。",
                    }
                )
            continue
        for stage, theme, label in labels:
            raw_hits[(stage, label)] += 1
            key = (msg["sender"], stage, label)
            if key in retained:
                duplicates.append(
                    {
                        "product": product,
                        "stage": stage,
                        "label": label,
                        "buyer": msg["sender"],
                        "date": msg["date"],
                        "time": msg["time"],
                        "duplicate_text": msg["text"],
                        "kept_text": retained[key]["text"],
                        "source_line": msg["source_line"],
                    }
                )
                continue
            retained[key] = msg
            buyers_by_label[(stage, theme, label)].add(msg["sender"])
            evidence_by_label[(stage, theme, label)].append(msg)
            detail.append(
                {
                    "product": product,
                    "stage": stage,
                    "theme": theme,
                    "label": label,
                    "buyer": msg["sender"],
                    "date": msg["date"],
                    "time": msg["time"],
                    "text": msg["text"],
                    "all_text": msg["text"],
                    "source_line": msg["source_line"],
                    "raw_hit": raw_hits[(stage, label)],
                }
            )
        for risk_label, _ in risk_items(msg["text"]):
            risk_rows.append(
                {
                    "product": product,
                    "risk_label": risk_label,
                    "buyer": msg["sender"],
                    "date": msg["date"],
                    "time": msg["time"],
                    "text": msg["text"],
                    "source_line": msg["source_line"],
                    "action": "安全/质量词需二次证据复核：确认是否为本人实际发生，记录档位、使用时长、批次和处理结果。",
                }
            )

    stats = []
    for (stage, theme, label), buyers in buyers_by_label.items():
        evidences = evidence_by_label[(stage, theme, label)]
        examples = "、".join(sorted(list(buyers))[:10])
        quotes = "｜".join([item["text"] for item in evidences[:5]])
        stats.append(
            {
                "product": product,
                "stage": stage,
                "theme": theme,
                "label": label,
                "count": len(buyers),
                "raw_hits": sum(1 for item in detail if item["stage"] == stage and item["label"] == label),
                "duplicates": sum(1 for item in duplicates if item["stage"] == stage and item["label"] == label),
                "priority": PRIORITY_MAP.get(label, "P1" if len(buyers) >= 5 else "P2"),
                "definition": label_definition(label),
                "buyers": examples,
                "quote": quotes,
                "action": ACTION_MAP.get(label, "沉淀为 V1 词条，后续观察人数、占比和客服消耗。"),
            }
        )

    dates = [msg["date"] for msg in valid]
    summary = {
        "product": product,
        "file": str(file),
        "raw_messages": len(raw),
        "valid_messages": len(valid),
        "conversations": max([msg["conversation_id"] for msg in raw] or [0]),
        "buyers": len(set(msg["sender"] for msg in valid)),
        "start_date": min(dates) if dates else "",
        "end_date": max(dates) if dates else "",
        "demand_count": len(detail),
        "presale_count": sum(1 for item in detail if item["stage"] == "售前"),
        "aftersale_count": sum(1 for item in detail if item["stage"] == "售后"),
        "duplicate_count": len(duplicates),
        "risk_count": len(risk_rows),
        "review_count": len(review_rows),
    }
    return {
        "summary": summary,
        "stats": stats,
        "detail": detail,
        "duplicates": duplicates,
        "risks": risk_rows,
        "reviews": review_rows,
        "valid_messages": valid,
    }


def label_definition(label):
    definitions = {
        "使用方法/教程": "已购买或收到后咨询开机、安装、调档、换头、卷发/直发具体操作。",
        "拆卸/更换配件": "咨询加热管、直发梳、卷筒、梳齿等配件拆装、更换和安装方法。",
        "不加热/温度低": "实际使用后反馈温度低、不加热、预热无效或达不到标称温度。",
        "不出雾/喷雾异常": "实际使用后反馈喷雾不出、雾量异常、水箱/加水相关问题。",
        "造型效果不好": "实际使用后反馈卷不出、夹不直、不定型、效果不明显或不持久。",
        "退款退货": "实际提出退款、退货、拒收、拦截、退换货或运费险相关售后处理。",
        "卷发/直发功能": "购买前咨询是否能卷发、直发、刘海、内扣、蓬松等造型。",
        "款式/型号区别": "购买前咨询不同型号、套餐、功能、价格和适用场景差异。",
        "温度/档位/预热": "购买前咨询温度范围、档位、预热时间和加热能力。",
        "是否伤发/发质适用": "购买前咨询是否伤发、是否适合自己的发质，以及安全顾虑。",
    }
    return definitions.get(label, f"{label}相关咨询或反馈。")


def sort_stats(stats):
    return sorted(stats, key=lambda item: (-item["count"], item["label"]))


def stage_stats(results, stage, product=None):
    rows = []
    for result in results:
        for item in result["stats"]:
            if item["stage"] == stage and (product is None or item["product"] == product):
                rows.append(item)
    return sort_stats(rows)


def aggregate_stats(results, stage):
    bucket = {}
    for item in stage_stats(results, stage):
        key = item["label"]
        if key not in bucket:
            bucket[key] = {**item, "count": 0, "raw_hits": 0, "duplicates": 0, "products": set(), "buyers_list": [], "quotes": []}
        bucket[key]["count"] += item["count"]
        bucket[key]["raw_hits"] += item["raw_hits"]
        bucket[key]["duplicates"] += item["duplicates"]
        bucket[key]["products"].add(item["product"])
        if item["buyers"]:
            bucket[key]["buyers_list"].append(item["buyers"])
        if item["quote"]:
            bucket[key]["quotes"].append(item["quote"])
    rows = []
    for row in bucket.values():
        row["product"] = "三款合计"
        row["buyers"] = "、".join(row["buyers_list"])[:200]
        row["quote"] = "｜".join(row["quotes"])[:300]
        row["products"] = "、".join(sorted(row["products"]))
        rows.append(row)
    return sort_stats(rows)


def style_workbook(wb):
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="微软雅黑", size=10, color=TEXT)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = thin_border


def title(ws, text, note=None, max_col=12):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    ws.cell(1, 1).value = text
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=RED)
    ws.cell(1, 1).font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    if note:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        ws.cell(2, 1).value = note
        ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 44


def header(ws, row_idx):
    for cell in ws[row_idx]:
        cell.fill = PatternFill("solid", fgColor="8D241F")
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, values):
    for idx, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_stat_sheet(ws, rows, title_text, total, days, buyer_count):
    title(ws, title_text, "本次为首次分析，无历史基准；本 Sheet 建立三款卷发棒 V1 需求词条。", 12)
    ws.append(["正式需求总数", total, "统计天数", days, "需求覆盖买家", buyer_count, None, None, None, None, None, None])
    ws.append([None] * 12)
    ws.append(["排名", "需求词条", "同人去重人数", "占比", "日均人数", "原始命中", "重复次数", "优先级", "严格定义", "涉及买家示例", "代表原话", "建议动作"])
    for idx, item in enumerate(rows, 1):
        ws.append([
            idx,
            item["label"],
            item["count"],
            item["count"] / total if total else 0,
            item["count"] / days if days else 0,
            item["raw_hits"],
            item["duplicates"],
            item["priority"],
            item["definition"],
            item.get("buyers", ""),
            item.get("quote", ""),
            item["action"],
        ])
    header(ws, 4)
    widths(ws, [8, 34, 15, 12, 12, 13, 12, 10, 52, 42, 68, 62])
    for row in range(5, ws.max_row + 1):
        ws.cell(row, 4).number_format = "0.0%"
        ws.cell(row, 5).number_format = "0.0"
    ws.freeze_panes = "A5"


PRODUCT_DISPLAY = {
    "856卷发棒": "856冷雾直板夹（直卷两用）",
    "366卷发棒": "366 32mm卷发棒",
    "五合一卷发棒": "五合一多头卷发棒",
}


def display_product(product):
    return PRODUCT_DISPLAY.get(product, product)


def product_prefix(value):
    text = str(value or "").strip()
    for prefix in ("856", "366", "五合一"):
        if text.startswith(prefix):
            return prefix
    return text


def infer_theme(label):
    theme_rules = [
        ("冷雾系统", ("冷雾", "出雾", "加水", "水箱", "漏水")),
        ("使用学习", ("使用方法", "教程", "换头", "拆装", "自动卷", "停留时间")),
        ("退换售后", ("退款", "退货", "运费险", "取件")),
        ("造型效果", ("造型", "卷径", "卷型", "直发", "卷发")),
        ("安全质量", ("伤发", "漏电", "冒烟", "烫伤", "异常", "二手", "耐用")),
        ("参数适配", ("温度", "档位", "功率", "发质", "发长", "尺寸", "重量", "电压")),
        ("交易与服务", ("优惠", "价保", "返现", "发货", "物流", "质保", "保修", "正品")),
        ("配件配置", ("配件", "赠品", "包装", "配置", "版本", "材质", "涂层")),
    ]
    for theme, keywords in theme_rules:
        if any(keyword in label for keyword in keywords):
            return theme
    return "其他需求"


def load_reference_taxonomy(path):
    """Load the confirmed V1 schema while deliberately ignoring historical counts."""
    if not path.exists():
        raise FileNotFoundError(f"找不到三款产品 V1 词条来源：{path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if "V1词条字典" not in workbook.sheetnames:
            raise ValueError("正式词条来源缺少 V1词条字典 Sheet")
        rows = []
        for values in workbook["V1词条字典"].iter_rows(min_row=2, values_only=True):
            label = str(values[0] or "").strip()
            if not label:
                continue
            product_names = [item.strip() for item in str(values[1] or "").split("、") if item.strip()]
            stage_text = str(values[2] or "")
            stages = [stage for stage in ("售前", "售后") if stage in stage_text]
            if not stages:
                continue
            priority = str(values[3] or "P2").upper()
            rows.append(
                {
                    "label": label,
                    "stage": stages[0],
                    "stages": stages,
                    "product_prefixes": sorted({product_prefix(item) for item in product_names}),
                    "theme": infer_theme(label),
                    "definition": str(values[5] or f"{label}相关的明确买家需求。").strip(),
                    "priority": priority if priority in ("P0", "P1", "P2") else "P2",
                    "action": str(values[6] or "持续观察并完善对应客服承接。").strip(),
                }
            )
        if len(rows) < 20:
            raise ValueError(f"正式 V1 词条读取不足：仅 {len(rows)} 个")
        return rows
    finally:
        workbook.close()


def result_by_prefix(results, prefix):
    for result in results:
        if result["summary"]["product"].startswith(prefix):
            return result
    raise ValueError(f"缺少 {prefix} 产品分析结果")


def clear_sheet_values(ws):
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def copy_row_format(ws, source_row, target_row, max_col):
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
    source_dimension = ws.row_dimensions[source_row]
    target_dimension = ws.row_dimensions[target_row]
    if source_dimension.height is not None:
        target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden


def write_rows(ws, start_row, rows, style_row=None, max_col=None):
    width = max_col or max((len(row) for row in rows), default=ws.max_column)
    template_max_row = ws.max_row
    for offset, values in enumerate(rows):
        row_idx = start_row + offset
        if style_row is not None and row_idx > template_max_row:
            copy_row_format(ws, style_row, row_idx, width)
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col_idx)
            if not isinstance(cell, MergedCell):
                cell.value = value


def product_stage_rows(result, stage, days):
    product = result["summary"]["product"]
    rows = stage_stats([result], stage, product)
    total = sum(item["count"] for item in rows)
    output = []
    for index, item in enumerate(rows, 1):
        output.append(
            [
                index,
                item["label"],
                item["count"],
                item["count"] / total if total else 0,
                item["count"] / days if days else 0,
                item["raw_hits"],
                item["duplicates"],
                item["priority"],
                item["definition"],
                item["buyers"],
                item["quote"],
                item["action"],
            ]
        )
    return rows, output


def top_stage_label(result, stage):
    rows = stage_stats([result], stage, result["summary"]["product"])
    if not rows:
        return "无"
    return f"{rows[0]['label']} {rows[0]['count']}人"


def write_template_overview(wb, results, summary_lines, days):
    ws = wb["分析总览"]
    clear_sheet_values(ws)
    starts = [result["summary"]["start_date"] for result in results if result["summary"]["start_date"]]
    ends = [result["summary"]["end_date"] for result in results if result["summary"]["end_date"]]
    start_date = min(starts) if starts else ""
    end_date = max(ends) if ends else ""
    ws["A1"] = f"丹橘卷发造型三款产品｜首次客服聊天需求基准（{start_date} 至 {end_date}）"
    ws["A2"] = (
        "首次分析，不套用其他商品历史词条。三款分别独立统计；同一商品＋买家＋阶段＋需求词条在本周期只计1次，"
        "一句话可拆多个需求。质量安全页保留风险候选，供人工二次校准。"
    )
    ordered = [
        result_by_prefix(results, "856"),
        result_by_prefix(results, "366"),
        result_by_prefix(results, "五合一"),
    ]
    for block, result in enumerate(ordered):
        summary = result["summary"]
        col = 1 + block * 4
        ws.cell(5, col).value = display_product(summary["product"])
        for offset, label in enumerate(["正式需求", "售前", "售后", "强风险"]):
            ws.cell(6, col + offset).value = label
        for offset, value in enumerate(
            [summary["demand_count"], summary["presale_count"], summary["aftersale_count"], summary["risk_count"]]
        ):
            ws.cell(7, col + offset).value = value
    write_rows(
        ws,
        9,
        [["产品", "有效买家", "正式需求", "售前", "售后", "售后占比", "重复剔除", None, "产品", "售前第一", "售后第一", "核心判断"]],
    )
    product_rows = []
    for result in ordered:
        summary = result["summary"]
        aftersale_ratio = summary["aftersale_count"] / summary["demand_count"] if summary["demand_count"] else 0
        product_rows.append(
            [
                display_product(summary["product"]),
                summary["buyers"],
                summary["demand_count"],
                summary["presale_count"],
                summary["aftersale_count"],
                aftersale_ratio,
                summary["duplicate_count"],
                None,
                summary["product"].replace("卷发棒", ""),
                top_stage_label(result, "售前"),
                top_stage_label(result, "售后"),
                "优先处理高频需求并持续复核风险候选",
            ]
        )
    write_rows(ws, 10, product_rows)
    ws["A15"] = "首次V1核心结论"
    conclusions = list(summary_lines) + [f"本期覆盖 {days} 个自然日；后续同口径报告可直接沿用本母版做趋势对比。"]
    for index, line in enumerate(conclusions[:6], 1):
        ws.cell(15 + index, 1).value = str(index)
        ws.cell(15 + index, 2).value = line


def write_template_product_sheets(wb, results, days):
    sheet_specs = [
        ("856", "856售前V1", "售前"),
        ("856", "856售后V1", "售后"),
        ("366", "366售前V1", "售前"),
        ("366", "366售后V1", "售后"),
        ("五合一", "五合一售前V1", "售前"),
        ("五合一", "五合一售后V1", "售后"),
    ]
    for prefix, sheet_name, stage in sheet_specs:
        ws = wb[sheet_name]
        result = result_by_prefix(results, prefix)
        summary = result["summary"]
        rows, data_rows = product_stage_rows(result, stage, days)
        total = sum(item["count"] for item in rows)
        clear_sheet_values(ws)
        ws["A1"] = f"{display_product(summary['product'])}｜首次{stage}V1"
        ws["A2"] = "本次为首次分析，无历史基准；本 Sheet 建立该产品 V1 需求词条。"
        write_rows(ws, 3, [["正式需求", total, "统计天数", days, "基准版本", "V1"]])
        write_rows(
            ws,
            4,
            [["排名", "需求词条", "去重人数", "阶段内占比", "日均", "原始命中", "重复次数", "优先级", "严格定义", "买家示例", "代表原话", "建议动作"]],
        )
        write_rows(ws, 5, data_rows, style_row=5, max_col=12)


def write_template_product_comparison(wb, results):
    ws = wb["三款产品对比"]
    clear_sheet_values(ws)
    ws["A1"] = "三款产品横向对比｜首次V1"
    write_rows(ws, 3, [["指标", "856冷雾直板夹", "366 32mm卷发棒", "五合一多头", "横向判断", "856占比", "366占比", "五合一占比", "建议", "备注"]])
    ordered = [result_by_prefix(results, "856"), result_by_prefix(results, "366"), result_by_prefix(results, "五合一")]
    summaries = [item["summary"] for item in ordered]

    def metric_row(name, key, judgement, suggestion, ratio_key=None):
        values = [summary[key] for summary in summaries]
        ratios = []
        for summary, value in zip(summaries, values):
            denominator = summary[ratio_key] if ratio_key else sum(values)
            ratios.append(value / denominator if denominator else 0)
        return [name, *values, judgement, *ratios, suggestion, "按同一统计口径"]

    rows = [
        metric_row("正式需求", "demand_count", "仅作规模参考", "结合曝光量和SKU内部占比判断"),
        metric_row("售前需求", "presale_count", "反映购买决策咨询", "优先优化详情页信息", "demand_count"),
        metric_row("售后需求", "aftersale_count", "反映实际使用承接压力", "优先优化教程和故障排查", "demand_count"),
        metric_row("重复剔除", "duplicate_count", "反映同人重复咨询", "关注高重复词条", "demand_count"),
        metric_row("风险候选", "risk_count", "需人工二次校准", "安全与质量问题单独升级", "demand_count"),
        metric_row("有效买家", "buyers", "反映样本覆盖", "下期沿用同口径比较"),
    ]
    write_rows(ws, 4, rows, style_row=4, max_col=10)


def write_template_risks(wb, results):
    ws = wb["质量安全风险"]
    clear_sheet_values(ws)
    ws["A1"] = "质量 / 安全 / 真实体验风险（AI 证据二次校准）"
    ws["A2"] = (
        "仅保留用户本人实际发生且通过独立证据二审的较强异常；购买前担忧、普通教程咨询、"
        "首次模糊提问和正常工作现象不进入本页。"
    )
    write_rows(ws, 4, [["ID", "产品", "风险类型", "风险词条", "买家", "日期", "时间", "代表原话", "全部证据", "建议动作", "优先级"]])
    rows = []
    for result in results:
        for item in result["risks"]:
            rows.append(
                [
                    len(rows) + 1,
                    display_product(item["product"]),
                    item["risk_label"],
                    item.get("demand_label", item["risk_label"]),
                    item["buyer"],
                    item["date"],
                    item["time"],
                    item["text"],
                    item["text"],
                    item["action"],
                    item.get("priority", "P1"),
                ]
            )
    if not rows:
        rows = [[1, "三款合计", "无", "无强风险候选", "—", "—", "—", "本期未命中", "本期未命中", "持续监测", "P2"]]
    write_rows(ws, 5, rows, style_row=5, max_col=11)


def write_template_cross_product(wb, results):
    ws = wb["跨产品共性"]
    clear_sheet_values(ws)
    ws["A1"] = "跨产品共性问题与内容优化"
    write_rows(ws, 3, [["共性问题", "影响产品", "主要差异", "建议动作", "优先级", None, "产品", "最应该先改什么", "原因", "预期减少的咨询"]])
    grouped = defaultdict(list)
    for result in results:
        for item in result["stats"]:
            grouped[item["label"]].append(item)
    common = []
    for label, items in grouped.items():
        products = sorted({display_product(item["product"]) for item in items})
        if len(products) < 2:
            continue
        common.append(
            [
                label,
                "、".join(products),
                "各SKU人数和占比不同",
                items[0]["action"],
                min((item["priority"] for item in items), default="P1"),
            ]
        )
    common.sort(key=lambda row: row[0])
    ordered = [result_by_prefix(results, "856"), result_by_prefix(results, "366"), result_by_prefix(results, "五合一")]
    focus = []
    for result in ordered:
        rows = stage_stats([result], "售后", result["summary"]["product"])
        top = rows[0] if rows else None
        focus.append(
            [
                display_product(result["summary"]["product"]),
                top["action"] if top else "持续监测",
                f"售后最高频为 {top['label']}" if top else "暂无售后词条",
                top["label"] if top else "暂无",
            ]
        )
    row_count = max(len(common), len(focus), 1)
    rows = []
    for index in range(row_count):
        left = common[index] if index < len(common) else [None] * 5
        right = focus[index] if index < len(focus) else [None] * 4
        rows.append([*left, None, *right])
    write_rows(ws, 4, rows, style_row=4, max_col=10)


def write_template_daily_trend(wb, results):
    ws = wb["每日趋势"]
    clear_sheet_values(ws)
    all_dates = [item["date"] for result in results for item in result["detail"]]
    start_date = min(all_dates) if all_dates else ""
    end_date = max(all_dates) if all_dates else ""
    ws["A1"] = f"每日需求趋势｜{start_date} 至 {end_date}"
    write_rows(ws, 3, [["日期", "产品", "需求买家", "售前日级需求", "售后日级需求", "强风险", "总日级需求", None, None, None]])
    rows = []
    ordered = [result_by_prefix(results, "856"), result_by_prefix(results, "366"), result_by_prefix(results, "五合一")]
    for result in ordered:
        product = result["summary"]["product"]
        daily = defaultdict(lambda: {"buyers": set(), "售前": 0, "售后": 0, "risk": 0})
        for item in result["detail"]:
            daily[item["date"]]["buyers"].add(item["buyer"])
            daily[item["date"]][item["stage"]] += 1
        for item in result["risks"]:
            daily[item["date"]]["risk"] += 1
        for date, values in sorted(daily.items()):
            rows.append(
                [
                    date,
                    display_product(product),
                    len(values["buyers"]),
                    values["售前"],
                    values["售后"],
                    values["risk"],
                    values["售前"] + values["售后"],
                    None,
                    None,
                    None,
                ]
            )
    write_rows(ws, 4, rows, style_row=4, max_col=10)


def write_template_detail_sheets(wb, results):
    ws = wb["逐条需求明细"]
    clear_sheet_values(ws)
    detail_rows = [["明细ID", "产品", "阶段", "需求词条", "买家", "日期", "时间", "代表原话", "同需求全部原话", "源日志行", "原始命中", "置信度"]]
    for result in results:
        for item in result["detail"]:
            detail_rows.append(
                [
                    len(detail_rows),
                    display_product(item["product"]),
                    item["stage"],
                    item["label"],
                    item["buyer"],
                    item["date"],
                    item["time"],
                    item["text"],
                    item["all_text"],
                    item["source_line"],
                    item["raw_hit"],
                    item.get("confidence", 0.0),
                ]
            )
    write_rows(ws, 1, detail_rows[:1])
    write_rows(ws, 2, detail_rows[1:], style_row=2, max_col=12)

    ws = wb["同人重复记录"]
    clear_sheet_values(ws)
    duplicate_rows = [["产品", "阶段", "需求词条", "买家", "日期", "时间", "重复原话", "源日志行", "首次保留原话", "剔除原因"]]
    for result in results:
        for item in result["duplicates"]:
            duplicate_rows.append(
                [
                    display_product(item["product"]),
                    item["stage"],
                    item["label"],
                    item["buyer"],
                    item["date"],
                    item["time"],
                    item.get("duplicate_text", item.get("text", "")),
                    item["source_line"],
                    item["kept_text"],
                    "同一商品+买家+阶段+词条，本周期只计1次",
                ]
            )
    write_rows(ws, 1, duplicate_rows[:1])
    write_rows(ws, 2, duplicate_rows[1:], style_row=2, max_col=10)

    ws = wb["待人工复核"]
    clear_sheet_values(ws)
    review_rows = [["产品", "买家", "日期", "时间", "原话", "源日志行", "复核原因", "状态"]]
    for result in results:
        for item in result["reviews"][:100]:
            review_rows.append(
                [
                    display_product(item["product"]),
                    item["sender"],
                    item["date"],
                    item["time"],
                    item["text"],
                    item["source_line"],
                    item["reason"],
                    "待人工确认",
                ]
            )
    if len(review_rows) == 1:
        review_rows.append(["三款合计", "—", "—", "—", "本期无需人工复核", "—", "—", "无需复核"])
    write_rows(ws, 1, review_rows[:1])
    write_rows(ws, 2, review_rows[1:], style_row=2, max_col=8)


def write_template_dictionary(wb, results, taxonomy):
    ws = wb["V1词条字典"]
    clear_sheet_values(ws)
    write_rows(ws, 1, [["V1需求词条", "适用产品", "本期阶段", "优先级", "三款合计去重人数", "严格定义", "建议动作"]])
    prefix_display = {
        "856": "856冷雾直板夹（直卷两用）",
        "366": "366 32mm卷发棒",
        "五合一": "五合一多头卷发棒",
    }
    grouped = {
        item["label"]: {
            "products": {prefix_display.get(prefix, prefix) for prefix in item.get("product_prefixes", [])},
            "stages": set(item.get("stages") or [item.get("stage", "售后")]),
            "count": 0,
            "priority": item["priority"],
            "definition": item["definition"],
            "action": item["action"],
        }
        for item in taxonomy
    }
    for result in results:
        for item in result["stats"]:
            row = grouped.setdefault(
                item["label"],
                {
                    "products": set(),
                    "stages": set(),
                    "count": 0,
                    "priority": item["priority"],
                    "definition": item["definition"],
                    "action": item["action"],
                },
            )
            row["products"].add(display_product(item["product"]))
            row["stages"].add(item["stage"])
            row["count"] += item["count"]
            row["priority"] = min(row["priority"], item["priority"])
    rows = []
    for label, item in sorted(grouped.items(), key=lambda pair: (-pair[1]["count"], pair[0])):
        rows.append(
            [
                label,
                "、".join(sorted(item["products"])),
                "、".join(sorted(item["stages"])),
                item["priority"],
                item["count"],
                item["definition"],
                item["action"],
            ]
        )
    write_rows(ws, 2, rows, style_row=2, max_col=7)


def write_template_raw_messages(wb, results):
    ws = wb["有效买家消息"]
    clear_sheet_values(ws)
    write_rows(ws, 1, [["产品", "买家", "日期", "时间", "阶段判断", "有效买家原话"]])
    rows = []
    for result in results:
        for item in result["valid_messages"]:
            stage = item.get("ai_stage", "未归类")
            rows.append(
                [
                    display_product(item["product"]),
                    item["sender"],
                    item["date"],
                    item["time"],
                    stage,
                    item["text"],
                ]
            )
    write_rows(ws, 2, rows, style_row=2, max_col=6)


def write_standard_template_report(output_xlsx, results, summary_lines, days, taxonomy):
    if not TEMPLATE_XLSX.exists():
        raise FileNotFoundError(f"找不到三款产品标准母版：{TEMPLATE_XLSX}")
    staging_xlsx = output_xlsx.with_name(f"{output_xlsx.stem}__staging.xlsx")
    wb = load_workbook(TEMPLATE_XLSX)
    expected_sheets = [
        "分析总览",
        "三款产品对比",
        "856售前V1",
        "856售后V1",
        "366售前V1",
        "366售后V1",
        "五合一售前V1",
        "五合一售后V1",
        "质量安全风险",
        "跨产品共性",
        "每日趋势",
        "逐条需求明细",
        "同人重复记录",
        "待人工复核",
        "V1词条字典",
        "口径说明",
        "有效买家消息",
    ]
    if wb.sheetnames != expected_sheets:
        raise ValueError("三款产品标准母版 Sheet 结构已变化，请先更新正式脚本映射。")
    write_template_overview(wb, results, summary_lines, days)
    write_template_product_comparison(wb, results)
    write_template_product_sheets(wb, results, days)
    write_template_risks(wb, results)
    write_template_cross_product(wb, results)
    write_template_daily_trend(wb, results)
    write_template_detail_sheets(wb, results)
    write_template_dictionary(wb, results, taxonomy)
    write_template_raw_messages(wb, results)
    ws = wb["口径说明"]
    starts = [result["summary"]["start_date"] for result in results if result["summary"]["start_date"]]
    ends = [result["summary"]["end_date"] for result in results if result["summary"]["end_date"]]
    if starts and ends:
        ws["B3"] = f"{min(starts)} 至 {max(ends)}"
        ws["C3"] = f"三份日志实际覆盖 {days} 天。"
    wb.save(staging_xlsx)
    write_values_into_exact_template(TEMPLATE_XLSX, staging_xlsx, output_xlsx)
    staging_xlsx.unlink(missing_ok=True)


def main():
    inputs = resolve_inputs()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    task_id = f"{timestamp}_丹橘三款卷发棒首次基准"
    output_dir = OUTPUT_ROOT / task_id
    output_dir.mkdir(parents=True, exist_ok=True)
    output_xlsx = output_dir / "丹橘三款卷发棒_20260824_首次客服聊天需求基准分析报告.xlsx"
    output_md = output_dir / "丹橘三款卷发棒_20260824_管理结论.md"
    manifest_path = output_dir / "result_manifest.json"

    entries = []
    for product, file in inputs:
        raw = parse_log(product, file)
        entries.append({"product": product, "raw": raw, "valid": valid_customer_messages(raw)})
    reference_taxonomy = load_reference_taxonomy(TAXONOMY_XLSX)
    results, taxonomy, summary_lines, llm_usage = analyze_entries_with_llm(
        entries,
        task_context="丹橘三款卷发造型产品客服聊天需求分析；三款产品分别统计并做跨产品比较。",
        reference_taxonomy=reference_taxonomy,
    )
    all_summary = [r["summary"] for r in results]
    total_days = len(set(msg["date"] for r in results for msg in r["valid_messages"])) or 1
    total_buyers = len(set(msg["sender"] for r in results for msg in r["valid_messages"]))
    total_demand = sum(s["demand_count"] for s in all_summary)
    total_presale = sum(s["presale_count"] for s in all_summary)
    total_aftersale = sum(s["aftersale_count"] for s in all_summary)
    total_risks = sum(s["risk_count"] for s in all_summary)
    total_duplicates = sum(s["duplicate_count"] for s in all_summary)

    presale_rows = aggregate_stats(results, "售前")
    aftersale_rows = aggregate_stats(results, "售后")
    top_presale = presale_rows[0] if presale_rows else None
    top_aftersale = aftersale_rows[0] if aftersale_rows else None

    write_standard_template_report(output_xlsx, results, summary_lines, total_days, taxonomy)
    output_md.write_text("# 丹橘三款卷发棒首次基准分析\n\n" + "\n".join(f"{i+1}. {line}" for i, line in enumerate(summary_lines)), encoding="utf-8")
    manifest = {
        "task_id": task_id,
        "product_name": "丹橘三款卷发棒",
        "analysis_type": "first_baseline",
        "status": "completed",
        "model": llm_usage["model"],
        "analysis_engine": "llm_api",
        "llm_usage": llm_usage,
        "taxonomy": taxonomy,
        "input_files": [str(file) for _, file in inputs],
        "output_files": {"excel": str(output_xlsx), "markdown": str(output_md), "manifest": str(manifest_path)},
        "summary": summary_lines,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


def add_compare_placeholder(wb, sheet_name, stage):
    ws = wb.create_sheet(sheet_name)
    title(ws, f"{stage}需求基准对比", "本次为首次分析，没有历史基准；本 Sheet 作为下一期对比模板保留。", 12)
    ws.append(["本期总需求", 0, "本期天数", 0, "基准总需求", 0, "基准天数", 0, None, None, None, None])
    ws.append([None] * 12)
    ws.append(["需求词条", "本期人数", "本期占比", "本期日均", "基准人数", "基准占比", "基准日均", "人数变化", "占比变化", "变化判断", "建议动作", "备注"])
    ws.append(["首次分析暂无基准", 0, 0, 0, 0, 0, 0, 0, 0, "首次基准", "下期沿用本期词条", "本期作为V1基准"])
    header(ws, 4)
    widths(ws, [34, 12, 12, 12, 12, 12, 12, 12, 12, 16, 62, 36])


def write_added_labels(wb, presale, aftersale):
    ws = wb.create_sheet("新增词条")
    ws.append(["阶段", "新增词条", "本期人数", "本期占比", "优先级", "严格定义", "建议动作"])
    for rows, stage in [(presale, "售前"), (aftersale, "售后")]:
        total = sum(item["count"] for item in rows) or 1
        for item in rows:
            ws.append([stage, item["label"], item["count"], item["count"] / total, item["priority"], item["definition"], item["action"]])
    header(ws, 1)
    widths(ws, [10, 36, 13, 13, 10, 60, 62])


def write_risks(wb, results):
    ws = wb.create_sheet("安全质量风险")
    ws.append(["风险ID", "产品", "风险词条", "买家", "日期", "时间", "代表原话", "源行", "二次复核状态", "建议动作"])
    idx = 1
    for result in results:
        for item in result["risks"]:
            ws.append([idx, item["product"], item["risk_label"], item["buyer"], item["date"], item["time"], item["text"], item["source_line"], "待人工二次确认", item["action"]])
            idx += 1
    if idx == 1:
        ws.append([1, "三款合计", "无强风险候选", "—", "—", "—", "本期未命中强风险词", "—", "无需复核", "持续监测"])
    header(ws, 1)
    widths(ws, [9, 16, 28, 24, 12, 10, 84, 14, 18, 62])


def write_detail(wb, results):
    ws = wb.create_sheet("逐条需求明细")
    ws.append(["明细ID", "产品", "阶段", "一级主题", "需求词条", "买家", "日期", "时间", "代表原话", "同需求全部原话", "源行", "原始命中"])
    idx = 1
    for result in results:
        for item in result["detail"]:
            ws.append([idx, item["product"], item["stage"], item["theme"], item["label"], item["buyer"], item["date"], item["time"], item["text"], item["all_text"], item["source_line"], item["raw_hit"]])
            idx += 1
    header(ws, 1)
    widths(ws, [9, 16, 10, 22, 35, 24, 12, 10, 68, 84, 14, 13])
    ws.freeze_panes = "A2"


def write_duplicates(wb, results):
    ws = wb.create_sheet("同人重复记录")
    ws.append(["产品", "阶段", "需求词条", "买家", "日期", "时间", "重复原话", "重复源行", "保留原话", "剔除原因"])
    for result in results:
        for item in result["duplicates"]:
            ws.append([item["product"], item["stage"], item["label"], item["buyer"], item["date"], item["time"], item["duplicate_text"], item["source_line"], item["kept_text"], "同一买家、同一阶段、同一需求词条，本周期只计1次"])
    header(ws, 1)
    widths(ws, [16, 10, 35, 24, 12, 10, 68, 14, 68, 54])
    ws.freeze_panes = "A2"


def write_reviews(wb, results):
    ws = wb.create_sheet("待人工复核")
    ws.append(["产品", "买家", "日期", "时间", "待复核原话", "源行", "复核原因", "状态"])
    rows = 0
    for result in results:
        for item in result["reviews"][:100]:
            ws.append([item["product"], item["sender"], item["date"], item["time"], item["text"], item["source_line"], item["reason"], "待人工确认"])
            rows += 1
    if rows == 0:
        ws.append(["三款合计", "—", "—", "—", "本期主要需求已完成规则归类", "—", "—", "无需复核"])
    header(ws, 1)
    widths(ws, [16, 24, 12, 10, 70, 14, 48, 14])


def write_labels(wb, presale, aftersale):
    ws = wb.create_sheet("词条与口径")
    ws.append(["需求词条", "本期阶段", "严格定义", "优先级", "售前人数", "售后人数", "基准状态", "建议动作"])
    by_label = {}
    for item in presale + aftersale:
        row = by_label.setdefault(item["label"], {"label": item["label"], "definition": item["definition"], "priority": item["priority"], "presale": 0, "aftersale": 0, "action": item["action"]})
        if item["stage"] == "售前":
            row["presale"] += item["count"]
        else:
            row["aftersale"] += item["count"]
    for row in sorted(by_label.values(), key=lambda x: (-(x["presale"] + x["aftersale"]), x["label"])):
        stage = "售前" if row["presale"] and not row["aftersale"] else "售后" if row["aftersale"] and not row["presale"] else "售前/售后"
        ws.append([row["label"], stage, row["definition"], row["priority"], row["presale"], row["aftersale"], "V1新增词条", row["action"]])
    header(ws, 1)
    widths(ws, [36, 16, 60, 10, 12, 12, 14, 64])


def write_baseline(wb, presale, aftersale):
    ws = wb.create_sheet("V1基准")
    ws.append(["阶段", "排名", "需求词条", "基准人数", "基准占比", "基准日均", "优先级", "严格定义", "建议动作"])
    for stage, rows in [("售前", presale), ("售后", aftersale)]:
        total = sum(item["count"] for item in rows) or 1
        for idx, item in enumerate(rows, 1):
            ws.append([stage, idx, item["label"], item["count"], item["count"] / total, "", item["priority"], item["definition"], item["action"]])
    header(ws, 1)
    widths(ws, [10, 8, 36, 13, 13, 12, 10, 60, 62])


def write_raw(wb, results):
    ws = wb.create_sheet("有效买家消息")
    ws.append(["产品", "会话序号", "买家", "阶段判断", "日期", "时间", "源行", "有效买家原话"])
    for result in results:
        for msg in result["valid_messages"]:
            labels = classify_message(msg["text"])
            stage = labels[0][0] if labels else "未归类"
            ws.append([msg["product"], msg["conversation_id"], msg["sender"], stage, msg["date"], msg["time"], msg["source_line"], msg["text"]])
    header(ws, 1)
    widths(ws, [16, 12, 24, 12, 12, 10, 12, 96])
    ws.freeze_panes = "A2"


if __name__ == "__main__":
    main()
