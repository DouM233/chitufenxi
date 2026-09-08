"""Inspect per-row fills/heights: mother template vs generated reports vs gold."""
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings("ignore")


def fill_summary(cell):
    fill = cell.fill
    if fill is None or fill.fill_type is None:
        return "-"
    color = getattr(fill.start_color, "rgb", None)
    theme = getattr(fill.start_color, "theme", None)
    if color and isinstance(color, str):
        return color[-6:]
    if theme is not None:
        return f"theme{theme}"
    return fill.fill_type or "?"


def survey(path, label, sheets, max_rows=14):
    print(f"\n===== {label} =====")
    wb = load_workbook(path, data_only=False, read_only=False)
    for name in sheets:
        if name not in wb.sheetnames:
            print(f"-- [{name}] 不存在")
            continue
        ws = wb[name]
        print(f"-- [{name}] max_row={ws.max_row}")
        shown = 0
        for row_idx in range(1, min(ws.max_row, 400) + 1):
            fills = [fill_summary(ws.cell(row_idx, col)) for col in range(1, 6)]
            ht = ws.row_dimensions[row_idx].height
            uniq = set(fills)
            if shown < max_rows or len(uniq) == 1 and fills[0] != "-":
                print(f"   r{row_idx:<4} ht={str(ht)[:6]:<7} fills(A-E)={fills}")
                shown += 1
            if row_idx > 60 and shown >= max_rows:
                break
    wb.close()


TEMPLATE = Path(r"D:\workspace\麦吉AI · 商品大脑\chitufenxi\templates\excel\单品客服需求分析标准母版.xlsx")
targets = sys.argv[1:]
SHEETS = ["P0风险明细", "逐条需求明细", "新增词条", "售前需求统计"]
survey(TEMPLATE, "母版模板", SHEETS, max_rows=20)
for t in targets:
    survey(t, Path(t).name, SHEETS, max_rows=14)
