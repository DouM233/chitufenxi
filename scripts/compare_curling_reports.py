from pathlib import Path

from openpyxl import load_workbook


FILES = [
    ("标准报告", Path(r"E:/新2/赤兔历史分析结果/丹橘卷发棒三款产品_20260726-0824_首次客服聊天需求基准分析报告.xlsx")),
    ("当前生成", Path(r"E:/新2/赤兔历史分析结果/03_最终报告/2026/08/20260826_155415_丹橘三款卷发棒首次基准/丹橘三款卷发棒_20260824_首次客服聊天需求基准分析报告.xlsx")),
]


def preview_row(ws, row, limit=12):
    return [ws.cell(row, col).value for col in range(1, min(ws.max_column, limit) + 1)]


def main():
    for label, file in FILES:
        wb = load_workbook(file, data_only=True)
        print(f"\n==== {label}: {file.name} ====")
        print("sheets:", wb.sheetnames)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"\n-- {sheet_name} rows={ws.max_row} cols={ws.max_column}")
            for row in range(1, min(ws.max_row, 6) + 1):
                print(row, preview_row(ws, row))


if __name__ == "__main__":
    main()
