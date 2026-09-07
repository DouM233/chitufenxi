"""Inspect E2E output quality: overview, top demands, risks."""
import sys
from pathlib import Path

from openpyxl import load_workbook

path = Path(sys.argv[1])
wb = load_workbook(path, data_only=True, read_only=True)

print(f"== {path.name} ==")
print("sheets:", wb.sheetnames)

overview = wb["分析总览"]
rows = [[c for c in r] for r in overview.iter_rows(values_only=True)]
print("\n-- 分析总览 关键行 --")
for r in rows[3:8]:
    print([v for v in r if v is not None][:12])

print("\n-- 售前需求统计 TOP5 --")
sheet = wb["售前需求统计"]
data = [[c for c in r] for r in sheet.iter_rows(values_only=True)]
for r in data[5:11]:
    if r and r[0] is not None:
        print(r[:5])

print("\n-- 售后需求统计 TOP5 --")
sheet = wb["售后需求统计"]
data = [[c for c in r] for r in sheet.iter_rows(values_only=True)]
for r in data[5:11]:
    if r and r[0] is not None:
        print(r[:5])

print("\n-- 安全质量风险 行数 --")
sheet = wb["安全质量风险"]
data = [[c for c in r] for r in sheet.iter_rows(values_only=True)]
print(len([r for r in data[1:] if r and any(v is not None for v in r)]), "条")
for r in data[1:6]:
    if r and any(v is not None for v in r):
        print("  ", r[1:4])

print("\n-- 核心结论 --")
for r in rows[22:30]:
    if r and r[1]:
        print("  •", str(r[1])[:120])
wb.close()
