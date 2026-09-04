from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FILES = [
    Path(r"E:/新2/赤兔历史分析结果/A17按摩枕_按运营原口径复算_20260726-0824_客服聊天分析报告.xlsx"),
    Path(r"E:/新2/赤兔历史分析结果/2449小气泡_20260812-0818_客服聊天需求分析报告.xlsx"),
]


def main():
    for file in FILES:
        print(f"\n==== {file.name} ====")
        wb = load_workbook(file, data_only=False)
        print("sheets:", wb.sheetnames)
        for ws in wb.worksheets:
            print(f"\n-- {ws.title} rows={ws.max_row} cols={ws.max_column} freeze={ws.freeze_panes}")
            print("merged:", [str(item) for item in list(ws.merged_cells.ranges)[:8]])
            widths = []
            for idx in range(1, min(ws.max_column, 10) + 1):
                letter = get_column_letter(idx)
                widths.append((letter, ws.column_dimensions[letter].width))
            print("widths:", widths)
            for row_idx in range(1, min(ws.max_row, 6) + 1):
                values = [ws.cell(row_idx, col_idx).value for col_idx in range(1, min(ws.max_column, 10) + 1)]
                print(row_idx, values)


if __name__ == "__main__":
    main()
