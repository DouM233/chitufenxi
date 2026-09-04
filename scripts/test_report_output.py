import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

import generate_report as report
from llm_analysis import aggregate_results, build_management_summary
from exact_excel_writer import write_values_into_exact_template
from generate_report import (
    TEMPLATE_XLSX,
    aggregate_stats,
    baseline_for_product,
    build_compare_rows,
    read_baseline_data,
    validate_stage_ranking,
    write_multi_product_sheets,
)


class ExactTemplateWriterTests(unittest.TestCase):
    def test_excel_illegal_control_characters_are_removed(self):
        workbook = Workbook()
        sheet = workbook.active
        report.set_cell(sheet, 1, 1, "正常中文\x04仍可写入")
        self.assertEqual("正常中文仍可写入", sheet["A1"].value)

    def test_template_values_do_not_leak_into_shorter_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            mother_path = root / "mother.xlsx"
            staging_path = root / "staging.xlsx"
            output_path = root / "output.xlsx"

            mother = Workbook()
            mother.active.title = "分析总览"
            mother["分析总览"]["A1"] = "旧标题"
            mother["分析总览"]["B14"] = "旧产品词条"
            mother.save(mother_path)

            staging = load_workbook(mother_path)
            staging["分析总览"]["A1"] = "新标题"
            staging["分析总览"]["B14"] = None
            staging.save(staging_path)

            write_values_into_exact_template(mother_path, staging_path, output_path)
            result = load_workbook(output_path, data_only=True)
            self.assertEqual(result["分析总览"]["A1"].value, "新标题")
            self.assertIsNone(result["分析总览"]["B14"].value)

    def test_multi_product_sheets_survive_exact_template_publish(self):
        def product_result(product, offset):
            return {
                "summary": {
                    "product": product,
                    "raw_messages": 100 + offset,
                    "valid_messages": 50 + offset,
                    "buyers": 20 + offset,
                    "demand_count": 12 + offset,
                    "presale_count": 7 + offset,
                    "aftersale_count": 5,
                    "duplicate_count": 2,
                    "risk_count": 1,
                    "review_count": 0,
                },
                "stats": [
                    {
                        "product": product,
                        "stage": "售前",
                        "theme": "选购",
                        "label": "款式区别",
                        "count": 7 + offset,
                        "raw_hits": 9 + offset,
                        "duplicates": 2,
                        "priority": "P1",
                        "definition": "购买前比较商品款式。",
                        "buyers": "买家A",
                        "quote": "三款有什么区别",
                        "action": "提供商品对比表。",
                    },
                    {
                        "product": product,
                        "stage": "售后",
                        "theme": "使用",
                        "label": "使用教程",
                        "count": 5,
                        "raw_hits": 6,
                        "duplicates": 1,
                        "priority": "P1",
                        "definition": "收货后询问使用方法。",
                        "buyers": "买家B",
                        "quote": "这个怎么使用",
                        "action": "提供使用视频。",
                    },
                ],
                "detail": [
                    {"product": product, "date": "2026-09-01"},
                    {"product": product, "date": "2026-09-02"},
                ],
                "valid_messages": [
                    {"date": "2026-09-01"},
                    {"date": "2026-09-02"},
                ],
            }

        results = [
            product_result("S-ZY-F4-3BU", 0),
            product_result("S-ZY-F4-4ABU", 1),
            product_result("S-ZY-F4-5WH", 2),
        ]
        baseline_products = {
            result["summary"]["product"]: {
                stage: {
                    "total": sum(item["count"] for item in result["stats"] if item["stage"] == stage),
                    "days": 2,
                    "rows": {
                        item["label"]: {
                            "count": item["count"],
                            "share": 1.0,
                            "daily": item["count"] / 2,
                            "action": item["action"],
                        }
                        for item in result["stats"]
                        if item["stage"] == stage
                    },
                }
                for stage in ("售前", "售后")
            }
            for result in results
        }
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            staging_path = root / "staging.xlsx"
            output_path = root / "output.xlsx"
            workbook = load_workbook(TEMPLATE_XLSX)
            created = write_multi_product_sheets(
                workbook,
                results,
                aggregate_stats(results, "售前"),
                aggregate_stats(results, "售后"),
                {"products": baseline_products, "taxonomy": [{"label": "使用教程"}]},
            )
            workbook.save(staging_path)
            write_values_into_exact_template(TEMPLATE_XLSX, staging_path, output_path)

            published = load_workbook(output_path, data_only=True)
            self.assertEqual(6, len(created))
            for title in created:
                self.assertIn(title, published.sheetnames)
            self.assertEqual("S-ZY-F4-3BU", published[created[0]]["A4"].value)
            self.assertEqual("S-ZY-F4-4ABU｜独立客服需求分析", published["S-ZY-F4-4ABU"]["A1"].value)
            values = [published["S-ZY-F4-4ABU"].cell(row, 1).value for row in range(1, published["S-ZY-F4-4ABU"].max_row + 1)]
            self.assertIn("售前基准对比", values)
            self.assertIn("售后基准对比", values)

    def test_reads_formal_multi_product_baseline(self):
        baseline = SCRIPT_DIR.parents[1] / "赤兔历史分析结果" / "苏泊尔足浴桶四SKU_20260724-0823_首次客服聊天需求基准分析报告.xlsx"
        if not baseline.exists():
            self.skipTest("formal multi-product baseline fixture is unavailable")
        data = read_baseline_data(baseline)
        self.assertEqual(57, len(data["taxonomy"]))
        self.assertEqual(4, len(data["products"]))
        self.assertEqual(2547, data["售前"]["total"])
        self.assertEqual(1547, data["售后"]["total"])
        self.assertEqual(31, data["售前"]["days"])
        self.assertGreaterEqual(len({item["theme"] for item in data["taxonomy"]}), 6)
        matches = [
            baseline_for_product(data, sku, 4)
            for sku in ("S-ZY-F4-5WH", "S-ZY-F4-3BU", "S-ZY-F4-4ABU", "S-ZY-F4-3AWH")
        ]
        self.assertTrue(all(matches))
        self.assertEqual(4, len({id(item) for item in matches}))

    def test_reads_standard_baseline_columns_by_header(self):
        baseline = SCRIPT_DIR.parents[1] / "赤兔历史分析结果" / "03_最终报告" / "2026" / "09" / "20260901_909化妆镜首次基准" / "909化妆镜_20260802-0831_首次客服聊天需求基准分析报告.xlsx"
        if not baseline.exists():
            self.skipTest("formal single-product baseline fixture is unavailable")
        data = read_baseline_data(baseline)
        item = next(row for row in data["taxonomy"] if row["label"] == "充电接口/电池续航参数")
        self.assertEqual("参数规格", item["theme"])
        self.assertEqual("P1", item["priority"])
        self.assertIn("充电接口类型", item["definition"])
        self.assertIn("详情页", item["action"])
        self.assertEqual("2026-08-02", data["start_date"])
        self.assertEqual("2026-08-31", data["end_date"])

    def test_reads_legacy_baseline_without_daily_column(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "旧版_20260723-0728_报告.xlsx"
            workbook = Workbook()
            overview = workbook.active
            overview.title = "分析总览"
            overview["A1"] = "旧版报告（2026-07-23 至 2026-07-28）"
            for stage, sheet_name in (("售前", "售前需求统计"), ("售后", "售后需求统计")):
                sheet = workbook.create_sheet(sheet_name)
                sheet.append(["旧版标题"])
                sheet.append([])
                sheet.append(["排名", "需求词条", "同人去重人数", "占比", "原始命中", "重复次数", "优先级", "严格定义", "买家", "原话", "建议动作"])
                sheet.append([1, f"{stage}词条", 12, 1, 20, 8, "P1", f"{stage}严格定义", "买家A", "原话", f"{stage}建议动作"])
            terms = workbook.create_sheet("词条与口径")
            terms.append(["词条、定义与口径"])
            terms.append([])
            terms.append(["需求词条", "适用阶段", "严格定义", "优先级", "建议动作"])
            terms.append(["售前词条", "售前", "售前严格定义", "P1", "售前建议动作"])
            terms.append(["售后词条", "售后", "售后严格定义", "P1", "售后建议动作"])
            workbook.save(path)

            data = read_baseline_data(path)
            item = data["售前"]["rows"]["售前词条"]
            self.assertEqual(6, data["售前"]["days"])
            self.assertEqual(2, item["daily"])
            self.assertEqual("P1", item["priority"])
            self.assertEqual("售前严格定义", item["definition"])
            self.assertEqual("售前建议动作", item["action"])
            taxonomy_item = next(row for row in data["taxonomy"] if row["label"] == "售前词条")
            self.assertEqual("售前建议动作", taxonomy_item["action"])


class BaselineComparisonTests(unittest.TestCase):
    def test_different_periods_use_daily_change_rate(self):
        current = [{"label": "使用教程", "count": 60, "action": "补充教程"}]
        baseline = {
            "total": 150,
            "days": 30,
            "rows": {
                "使用教程": {
                    "count": 90,
                    "share": 0.6,
                    "daily": 3,
                    "action": "补充教程",
                }
            },
        }
        row = build_compare_rows(current, total=100, days=10, baseline_stage=baseline)[0]
        self.assertEqual(1.0, row[9])
        self.assertEqual("日均明显上升", row[10])

    def test_equal_active_days_do_not_imply_same_period(self):
        current = [{"label": "使用教程", "count": 15, "action": "补充教程"}]
        baseline = {
            "total": 100,
            "days": 10,
            "rows": {"使用教程": {"count": 10, "share": 0.1, "daily": 1, "action": "补充教程"}},
        }
        row = build_compare_rows(current, total=150, days=10, baseline_stage=baseline, same_period=False)[0]
        self.assertEqual("日均明显上升", row[10])

    def test_baseline_terms_are_reused_and_only_true_additions_are_listed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            baseline_path = root / "baseline.xlsx"
            chat_path = root / "本期商品聊天记录.log"
            output_path = root / "output.xlsx"

            baseline = Workbook()
            baseline.active.title = "售前需求统计"
            baseline.create_sheet("售后需求统计")
            baseline.create_sheet("词条与口径")
            for stage, sheet_name, label in (("售前", "售前需求统计", "旧售前词条"), ("售后", "售后需求统计", "旧售后词条")):
                sheet = baseline[sheet_name]
                sheet.cell(4, 1, "正式需求总数")
                sheet.cell(4, 2, 10)
                sheet.cell(4, 3, "统计天数")
                sheet.cell(4, 4, 10)
                sheet.append([])
                sheet.cell(6, 1, "排名")
                sheet.cell(6, 2, "需求词条")
                sheet.cell(7, 1, 1)
                sheet.cell(7, 2, label)
                sheet.cell(7, 3, 10)
                sheet.cell(7, 4, 1)
                sheet.cell(7, 5, 1)
                sheet.cell(7, 8, "P1")
                sheet.cell(7, 9, f"{label}定义")
                sheet.cell(7, 12, f"{label}动作")
            terms = baseline["词条与口径"]
            terms.append(["需求词条", "本期阶段", "严格定义", "优先级", "售前人数", "售后人数", "生成状态", "建议动作"])
            terms.append(["旧售前词条", "售前", "旧售前词条定义", "P1", 10, 0, "V1", "旧售前动作"])
            terms.append(["旧售后词条", "售后", "旧售后词条定义", "P1", 0, 10, "V1", "旧售后动作"])
            baseline.save(baseline_path)
            chat_path.write_text("以下为一通会话\n买家A 2026-09-01 10:00:00\n本期问题\n会话结束", encoding="utf-8")

            taxonomy = [
                {"label": "旧售前词条", "stage": "售前", "stages": ["售前"], "product_prefixes": [], "theme": "购买", "definition": "旧售前词条定义", "priority": "P1", "action": "旧售前动作"},
                {"label": "旧售后词条", "stage": "售后", "stages": ["售后"], "product_prefixes": [], "theme": "使用", "definition": "旧售后词条定义", "priority": "P1", "action": "旧售后动作"},
                {"label": "本期新增词条", "stage": "售后", "stages": ["售后"], "product_prefixes": [], "theme": "使用", "definition": "本期新增定义", "priority": "P1", "action": "本期新增动作"},
            ]
            valid = {"product": "本期商品", "sender": "买家A", "date": "2026-09-01", "time": "10:00:00", "conversation_id": 1, "source_line": 2, "text": "本期问题"}
            stats = [
                {"product": "本期商品", "stage": "售前", "theme": "购买", "label": "旧售前词条", "count": 1, "raw_hits": 1, "duplicates": 0, "priority": "P1", "definition": "旧售前词条定义", "buyers": "买家A", "quote": "本期问题", "action": "旧售前动作"},
                {"product": "本期商品", "stage": "售后", "theme": "使用", "label": "本期新增词条", "count": 1, "raw_hits": 1, "duplicates": 0, "priority": "P1", "definition": "本期新增定义", "buyers": "买家A", "quote": "本期问题", "action": "本期新增动作"},
            ]
            result = {
                "summary": {"product": "本期商品", "raw_messages": 1, "valid_messages": 1, "conversations": 1, "buyers": 1, "start_date": "2026-09-01", "end_date": "2026-09-01", "demand_count": 2, "presale_count": 1, "aftersale_count": 1, "duplicate_count": 0, "risk_count": 0, "review_count": 0},
                "stats": stats,
                "detail": [{**valid, "stage": "售前", "label": "旧售前词条", "theme": "购买", "all_text": "本期问题", "raw_hit": 1}],
                "duplicates": [],
                "risks": [],
                "reviews": [],
                "valid_messages": [valid],
            }
            data = {
                "product_name": "本期商品",
                "analysis_type": "baseline_compare",
                "user_message": "本期商品与上期基准对比",
                "chat_files": [{"path": str(chat_path), "name": chat_path.name}],
                "baseline_files": [{"path": str(baseline_path), "name": baseline_path.name}],
            }
            with patch.object(
                report,
                "analyze_entries_with_llm",
                return_value=([result], taxonomy, ["测试结论"], {"analysis_complete": True, "expected_messages": 1, "analyzed_messages": 1, "degraded_messages": 0}),
            ) as analyzer:
                report.write_workbook(data, output_path)

            kwargs = analyzer.call_args.kwargs
            self.assertEqual("旧售前词条", kwargs["reference_taxonomy"][0]["label"])
            self.assertEqual(10, kwargs["baseline_metrics"]["售前"]["total"])
            output = load_workbook(output_path, data_only=True)
            self.assertEqual("本期新增词条", output["新增词条"]["B2"].value)
            self.assertEqual("旧售前词条", output["上期基准"]["C2"].value)
            terms_values = [output["词条与口径"].cell(row, 1).value for row in range(2, 5)]
            self.assertIn("旧售后词条", terms_values)
            self.assertEqual("购买", output["词条与口径"]["C2"].value)
            self.assertEqual("旧售前词条定义", output["词条与口径"]["D2"].value)
            self.assertEqual("P1", output["词条与口径"]["E2"].value)
            self.assertEqual(2, output["有效买家消息"].max_row)


class EvidenceTraceTests(unittest.TestCase):
    def test_retained_detail_collects_all_duplicate_quotes(self):
        messages = [
            {"_llm_id": "m1", "product": "商品A", "sender": "买家A", "date": "2026-09-01", "time": "10:00:00", "source_line": 1, "text": "怎么安装"},
            {"_llm_id": "m2", "product": "商品A", "sender": "买家A", "date": "2026-09-01", "time": "10:01:00", "source_line": 2, "text": "安装视频有吗"},
        ]
        entries = [{"product": "商品A", "raw": messages, "valid": messages}]
        taxonomy = [{"label": "安装教程", "stage": "售后", "stages": ["售后"], "theme": "使用", "definition": "安装方法", "priority": "P1", "action": "提供视频"}]
        demands = {
            "m1": [{"label": "安装教程", "stage": "售后", "confidence": 1}],
            "m2": [{"label": "安装教程", "stage": "售后", "confidence": 1}],
        }
        result = aggregate_results(entries, taxonomy, demands, {}, {})[0]
        self.assertEqual(1, len(result["detail"]))
        self.assertEqual(1, len(result["duplicates"]))
        self.assertEqual("怎么安装｜安装视频有吗", result["detail"][0]["all_text"])
        self.assertEqual(2, result["detail"][0]["raw_hit"])


class ManagementSummaryTests(unittest.TestCase):
    def test_daily_and_share_values_are_preformatted_for_the_model(self):
        class Client:
            def __init__(self):
                self.prompt = ""
                self._usage_lock = threading.Lock()
                self.attempts = 0
                self.max_attempts = None

            def complete_json(self, system_prompt, user_prompt, max_tokens):
                self.prompt = user_prompt
                return {"summary": ["一", "二", "三"]}

        client = Client()
        results = [{
            "summary": {"product": "商品A", "active_days": 2},
            "stats": [{"stage": "售前", "label": "参数咨询", "count": 3, "priority": "P1", "action": "补参数"}],
            "risks": [],
        }]
        baseline = {
            "period": {"start_date": "2026-08-01", "end_date": "2026-08-31"},
            "售前": {"total": 10, "days": 5, "rows": [{"label": "参数咨询", "count": 1, "share": 0.1, "daily": 0.2}]},
            "售后": {"total": 0, "days": 5, "rows": []},
        }
        build_management_summary(client, results, [], [], baseline_metrics=baseline)
        self.assertIn('"daily": "1.50"', client.prompt)
        self.assertIn('"share": "10.0%"', client.prompt)


class RankingTests(unittest.TestCase):
    def test_accepts_count_descending_and_label_tie_break(self):
        validate_stage_ranking(
            [
                {"label": "安装教程", "count": 10},
                {"label": "卡扣异常", "count": 5},
                {"label": "破损", "count": 5},
            ],
            "售后",
        )

    def test_rejects_unsorted_rows(self):
        with self.assertRaises(RuntimeError):
            validate_stage_ranking(
                [
                    {"label": "低频", "count": 2},
                    {"label": "高频", "count": 9},
                ],
                "售前",
            )


if __name__ == "__main__":
    unittest.main()
