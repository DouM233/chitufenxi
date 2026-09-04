import hashlib
import json
import sys
import warnings
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


template_path = Path(sys.argv[1])
output_path = Path(sys.argv[2])


def sha256(data):
    return hashlib.sha256(data).hexdigest()


with zipfile.ZipFile(template_path) as template_zip, zipfile.ZipFile(output_path) as output_zip:
    template_bad = template_zip.testzip()
    output_bad = output_zip.testzip()
    template_names = set(template_zip.namelist())
    output_names = set(output_zip.namelist())
    non_sheet_names = sorted(name for name in template_names if not name.startswith("xl/worksheets/"))
    changed_non_sheet = [
        name
        for name in non_sheet_names
        if name not in output_names or template_zip.read(name) != output_zip.read(name)
    ]
    xml_errors = []
    for name in sorted(output_names):
        if not name.endswith((".xml", ".rels")):
            continue
        try:
            ET.fromstring(output_zip.read(name))
        except Exception as exc:
            xml_errors.append({"part": name, "error": str(exc)})
    template_styles = {
        cell.attrib.get("s", "0")
        for name in template_names
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        for cell in ET.fromstring(template_zip.read(name)).iter()
        if cell.tag.endswith("}c")
    }
    output_styles = {
        cell.attrib.get("s", "0")
        for name in output_names
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        for cell in ET.fromstring(output_zip.read(name)).iter()
        if cell.tag.endswith("}c")
    }

with warnings.catch_warnings():
    warnings.simplefilter("ignore")
    workbook = load_workbook(output_path, data_only=False, read_only=False)

expected_sheets = [
    "分析总览",
    "三款产品对比",
    "856售前V1",
    "856售后V1",
    "366售前V1",
    "366售后V1",
    "五合一售前V1",
    "五合一售后V1",
    "质量安全风险",
    "跨产品共性",
    "每日趋势",
    "逐条需求明细",
    "同人重复记录",
    "待人工复核",
    "V1词条字典",
    "口径说明",
    "有效买家消息",
]

error_values = []
placeholder_hits = []
sheet_stats = []
for sheet in workbook.worksheets:
    non_empty = 0
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if value not in (None, ""):
                non_empty += 1
            if isinstance(value, str):
                if value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                    error_values.append(f"{sheet.title}!{cell.coordinate}={value}")
                if "待真实分析写入" in value:
                    placeholder_hits.append(f"{sheet.title}!{cell.coordinate}")
    sheet_stats.append(
        {
            "name": sheet.title,
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "non_empty": non_empty,
            "charts": len(sheet._charts),
            "merged_ranges": len(sheet.merged_cells.ranges),
        }
    )

overview_values = []
overview = workbook["分析总览"]
for row in overview.iter_rows(min_row=1, max_row=min(18, overview.max_row), min_col=1, max_col=min(12, overview.max_column)):
    values = [cell.value for cell in row]
    if any(value not in (None, "") for value in values):
        overview_values.append(values)

report = {
    "template": str(template_path),
    "output": str(output_path),
    "template_zip_bad_part": template_bad,
    "output_zip_bad_part": output_bad,
    "zip_part_sets_equal": template_names == output_names,
    "changed_non_worksheet_parts": changed_non_sheet,
    "all_xml_parts_valid": not xml_errors,
    "xml_errors": xml_errors,
    "styles_hash_equal": sha256(zipfile.ZipFile(template_path).read("xl/styles.xml"))
    == sha256(zipfile.ZipFile(output_path).read("xl/styles.xml")),
    "output_style_ids_subset_of_template": output_styles.issubset(template_styles),
    "sheet_names_exact": workbook.sheetnames == expected_sheets,
    "sheet_count": len(workbook.sheetnames),
    "chart_count": sum(len(sheet._charts) for sheet in workbook.worksheets),
    "merged_range_count": sum(len(sheet.merged_cells.ranges) for sheet in workbook.worksheets),
    "formula_or_value_errors": error_values,
    "placeholder_hits": placeholder_hits,
    "sheet_stats": sheet_stats,
    "overview_preview": overview_values,
}
print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
